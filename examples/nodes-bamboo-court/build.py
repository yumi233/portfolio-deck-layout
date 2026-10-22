"""Worked example: 有节·竹院 茶家具设计报告 (42 slides, 1200x675pt) -> 51-slide bilingual deck (960x540pt).

  python -I build.py SOURCE.pptx OUT.pptx WORKDIR

Everything deck-specific lives here (shape names, colour maps, slide order, per-slide layouts); the generic
machinery is scripts/pptx_kit.py and scripts/assets.py. Use it as a pattern, not as code to run on another deck:
read the new deck with scripts/inventory.py first and write the same kind of plan for it.
"""
import sys, os, io, hashlib, json
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'scripts'))
sys.path.insert(0, HERE)
from PIL import Image
from pptx import Presentation
import pptx_kit as K
import assets as AS
import fonts as FN
from en_text import EN, ROW_FLOW, TAGLINE_EN, FORM_STUDY_EN, CASE16

SRC, OUT, WORK = sys.argv[1:4]
MEDIA, A = os.path.join(WORK, 'media'), os.path.join(WORK, 'assets')
os.makedirs(MEDIA, exist_ok=True); os.makedirs(A, exist_ok=True)

cn_font, _ = FN.find(FN.CN); en_font, en_file = FN.find(FN.EN)
K.configure(en_font=en_font, cn_font=cn_font, en_font_file=en_file)

prs = Presentation(SRC)
prs.slide_width, prs.slide_height = K.E(960), K.E(540)
S = {i + 1: s for i, s in enumerate(prs.slides)}

# ------------------------------------------------------------------ 1. images: extract once, build derived assets
def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        elif s.shape_type == 13:
            yield s
for sl in prs.slides:
    for s in walk(sl.shapes):
        p = os.path.join(MEDIA, s.image.sha1[:8] + '.bin')
        if not os.path.exists(p):
            open(p, 'wb').write(s.image.blob)
img = lambda h: Image.open(os.path.join(MEDIA, h + '.bin'))

RENDERS = {'5daefaf6': 'host', '691ab745': 'guest', 'e3b9ab7d': 'table', 'e95188a2': 'twin', '3d1eb110': 'shelf', '0f72668c': 'lamp'}
for h, k in RENDERS.items():
    cut = AS.trim(AS.white_to_alpha(img(h)), thr=10); cut.save(f'{A}/cut_{k}.png')
    AS.motion_blur(cut).save(f'{A}/blur_{k}.png')
CUT = {h: f'{A}/cut_{k}.png' for h, k in RENDERS.items()}
for h in ['603d77fc', '8d7f3f5b', 'ac8d0326', '69ec03cf', '9a157f22', 'efe0c995']:      # S4 timeline figures
    AS.archival_gray(img(h)).save(f'{A}/gray_{h}.png')
AS.trim(img('26b1a91d').convert('RGBA'), thr=10, pad=6).save(f'{A}/line_actions.png')
hands = AS.trim(img('3aa9f4ef').convert('RGBA'), thr=10, pad=6)
AS.lineart(hands.crop((int(hands.width * 0.5), 0, hands.width, hands.height)), gray=138).save(f'{A}/toc_hands.png')
AS.trim(img('1c5f7226').convert('RGBA'), thr=40, pad=4).save(f'{A}/sketch_table.png')
AS.trim(img('91231817').convert('RGBA'), thr=40, pad=4).save(f'{A}/line_journey.png')
c1, c2 = AS.knockout(img('608752ff')), AS.knockout(img('d00cffb0'))
c1 = c1.resize((round(c1.width * 600 / c1.height), 600), Image.LANCZOS)
c2 = c2.resize((round(c2.width * 516 / c2.height), 516), Image.LANCZOS)
can = Image.new('RGBA', (c1.width + 150 + c2.width, 600), (255, 255, 255, 0))
can.alpha_composite(c1, (0, 0)); can.alpha_composite(c2, (c1.width + 150, 600 - c2.height)); can.save(f'{A}/chairs_gray.png')
for name, spec in {
    'cover_family.png': {'canvas': [480, 320], 'items': [
        {'src': 'cut_shelf.png', 'x': 8, 'y': 6, 'h': 170}, {'src': 'cut_lamp.png', 'x': 436, 'y': 0, 'h': 150},
        {'src': 'cut_table.png', 'x': 118, 'y': 52, 'w': 250}, {'src': 'cut_host.png', 'x': 22, 'y': 168, 'h': 118},
        {'src': 'cut_twin.png', 'x': 150, 'y': 196, 'w': 196}, {'src': 'cut_guest.png', 'x': 372, 'y': 170, 'h': 112}]},
    'lineup.png': {'canvas': [560, 112], 'items': [
        {'src': 'cut_shelf.png', 'x': 0, 'y': 0, 'h': 110}, {'src': 'cut_table.png', 'x': 88, 'y': 38, 'w': 150},
        {'src': 'cut_host.png', 'x': 250, 'y': 40, 'h': 70}, {'src': 'cut_guest.png', 'x': 318, 'y': 42, 'h': 68},
        {'src': 'cut_twin.png', 'x': 390, 'y': 50, 'w': 118}, {'src': 'cut_lamp.png', 'x': 520, 'y': 0, 'h': 110}]},
}.items():
    AS.compose(spec, A).save(f'{A}/{name}')

# ------------------------------------------------------------------ 2. source-deck colour language -> layout palette
def chapter(i):
    return 0 if i <= 2 else 1 if i <= 10 else 2 if i <= 17 else 3 if i <= 26 else 4 if i <= 40 else 5
RESEARCH = (1, 2)                                     # tea red allowed only in these chapters
TEXT = {'20342F': K.TITLE, '182C27': K.TITLE, '1B302A': K.TITLE, '0D1C1A': K.TITLE,
        '4A514D': K.BODY, '515C55': K.BODY, '7A817B': K.AUX, '727171': K.AUX}
LINE = {'ADB7AD': K.HAIR, 'C8CBC5': K.HAIR, 'C7D0C6': K.HAIR, 'C8D1C8': K.HAIR, 'AFB4AE': K.OUTL, '96AAA7': K.OUTL,
        '20342F': K.MIDL, '0D1C1A': K.MIDL, '1B302A': K.MIDL, '698A77': K.MIDL, 'BCC8BE': 'FFFFFF'}
FILL = {'EDF0E9': 'FFFFFF', 'EEF0E9': 'FFFFFF', 'E3EAE2': 'FFFFFF', 'F7F7F4': 'FFFFFF', 'F6F6F2': 'FFFFFF',
        'DFE4DD': 'E6E6E4', 'E0E5DE': 'E6E6E4', 'BCC8BE': 'C9CACA', 'C7D0C6': 'DCDCDA'}
ACCENTS = {'718779', '8A715A', '82654E', '73929C', '698A77', '604C3F', 'A17B59', '638A99', '628F7A'}
OVERRIDES = {
    4: {'text': {a: K.BODY for a in ('718779', '8A715A', '82654E', '73929C')}},
    5: {'fill': {'628F7A': K.RED, 'A17B59': 'A9686A', '638A99': 'C49A9C'},
        'line': {'628F7A': K.RED, 'A17B59': 'A9686A', '638A99': 'C49A9C'}},
    6: {'fill': {'A9BEA1': 'C79A9C', 'A5C1BA': 'DDBFC0', 'C9BA99': 'E8D6D7', 'B8BDA6': 'B98284', 'B49B57': K.RED,
                 'A96E50': K.RED, '668C8E': K.RED, '678A57': K.RED, '6E7152': K.RED, '8AA8A0': K.RED},
        'line': {'A9BEA1': 'FFFFFF', 'A5C1BA': 'FFFFFF', 'C9BA99': 'FFFFFF', 'B8BDA6': 'FFFFFF', 'C7D0C6': K.HAIR,
                 'B49B57': K.RED, 'A96E50': K.RED}},
    14: {'fill': {'8A715A': K.OUTL, '718779': K.OUTL, '73929C': K.OUTL, '20342F': K.RED},
         'text': {'8A715A': K.AUX, '718779': K.AUX, '73929C': K.AUX}},
}
def base(sl, i):
    K.drop_chrome(sl, names=('Image 2', 'background-softener', 'v20-running', 'v20-page'), texts=('NODES / BAMBOO COURT',))
    K.scale_geometry(sl, 0.8)
    ov = dict(OVERRIDES.get(i, {}))
    if chapter(i) in RESEARCH:          # dark green dots/fills become tea red in research chapters
        ov.setdefault('fill', {}); [ov['fill'].setdefault(c, K.RED) for c in ('20342F', '0D1C1A', '1B302A')]
    else:
        ov.setdefault('fill', {}); [ov['fill'].setdefault(c, K.AUX) for c in ('20342F', '0D1C1A', '1B302A')]
    K.recolor(sl, chapter(i) in RESEARCH, TEXT, LINE, FILL, ACCENTS, ov, tag=f'S{i}')
    K.retype(sl)
    K.restyle_tables(sl)
    K.white_bg(sl)

def take_header(sl):
    cn = K.take_text(sl, 'v20-title')
    en = K.en_case(K.take_text(sl, 'v20-subtitle'))
    return cn, en

# ------------------------------------------------------------------ 3. per-slide layouts that need more than vfit
def layout_timeline(sl):
    """S4 zig-zag timeline: entries alternate above/below the spine; one short EN line under each CN copy."""
    tx, tw = [60, 210, 360, 510, 660, 807], [135, 140, 135, 140, 135, 93]
    lead = K.find(sl, 'history-lead')
    K.add_en(sl, lead, EN[(4, 'history-lead')], width=900 - K.bbox(lead)[0]); K.move_to(lead, y=72)
    def stack(k, y):
        for nm, gap in (('date', 6), ('name', 8), ('copy', 8), ('meta', 0)):
            el = K.find(sl, f'history-{nm}-{k}')
            if nm == 'copy':
                K.add_en(sl, el, EN[(4, f'history-copy-{k}')], width=tw[k])
            else:
                K.fit_height(el)
            K.move_to(el, tx[k], y); y += K.bbox(el)[3] + gap
        return y
    imgs = {k: K.find(sl, f'history-image-{k}') for k in range(6)}
    top = K.bbox(lead)[1] + K.bbox(lead)[3] + 12
    row1 = max(stack(k, top) for k in (0, 2, 4))
    spine_y = max(row1, top + max(K.bbox(imgs[k])[3] for k in (1, 3, 5))) + 14
    for k in (1, 3, 5):
        K.move_to(imgs[k], y=spine_y - 14 - K.bbox(imgs[k])[3])
    sp = K.find(sl, 'history-spine'); K.move_to(sp, y=spine_y)
    for k in range(6):
        t = K.find(sl, f'history-tick-{k}'); K.move_to(t, y=spine_y + K.bbox(sp)[3] / 2 - K.bbox(t)[3] / 2)
    row2 = max(stack(k, spine_y + 16) for k in (1, 3, 5))
    for k in (0, 2, 4):
        K.move_to(imgs[k], y=spine_y + 14)
    end = max(row2, max(K.bbox(imgs[k])[1] + K.bbox(imgs[k])[3] for k in (0, 2, 4))) + 14
    K.move_to(K.find(sl, 'history-end-rule'), y=end)
    focus = K.find(sl, 'history-design-focus')
    K.add_en(sl, focus, EN[(4, 'history-design-focus')], width=500); K.move_to(focus, y=end + 10)

def layout_process(sl):
    """S24: five process rows sized to their CN/EN copy, centred as a block."""
    rows = []
    for k in range(5):
        cp = K.find(sl, f'process-copy-{k}')
        K.add_en(sl, cp, EN[(24, f'process-copy-{k}')], width=232)
        rows.append(max(K.bbox(cp)[3], K.bbox(K.find(sl, f'process-image-{k}'))[3]))
    y = 74 + max(0, (416 - sum(rows) - 14 * 4) / 2)
    for k, rh in enumerate(rows):
        for nm in ('image', 'no', 'name', 'rule', 'copy'):
            el = K.find(sl, f'process-{nm}-{k}'); K.move_to(el, y=y + rh / 2 - K.bbox(el)[3] / 2)
        y += rh + 14

CUSTOM = {4: layout_timeline, 24: layout_process}
ANCHORS = {23: [(f'material-no-{k}', f'material-name-{k}', 1) for k in range(4)],
           12: [(f'furniture-point-{k}', 'furniture-spine', 'c') for k in range(3)],
           21: [('user-context-title', 'user-context-image', ('below', 18)), ('user-context-copy', 'user-context-title', ('below', 8))]}

def sketch_page(sl, i):
    """H: sketch left; kicker / tag / copy (+EN) stacked on the right; product name moves to the header."""
    base(sl, i)
    name = K.take_text(sl, 'product-form-name'); en = K.en_case(K.take_text(sl, 'product-form-en'))
    K.header(sl, name, en)
    K.place(K.find(sl, 'product-form-plate'), 56, 92, 552, 368)
    kick, tag, cp = (K.find(sl, n) for n in ('product-form-kicker', 'product-form-tag', 'product-form-copy'))
    for el, sz in ((kick, 10), (tag, 14), (cp, 12)):
        K.left_align(el); K.set_size(el, sz)
    K.add_en(sl, kick, ('inline', FORM_STUDY_EN))
    K.add_en(sl, tag, EN[(i, 'product-form-tag')], width=260)
    K.add_en(sl, cp, EN[(i, 'product-form-copy')], width=260)
    K.text_block(sl, [kick, tag, cp], 650, 276, [0, 8, 16])

def render_page(sl, i):
    """I: cut-out render left, key / copy / materials (+EN) right."""
    base(sl, i)
    K.header(sl, *take_header(sl))
    hero = K.find(sl, 'product-final')
    K.swap_cutout(sl, hero, CUT[K.pic_sha(sl, hero)[0]], box=(64, 92, 496, 384))
    els, gaps = [], []
    for nm, sz, gap in (('product-final-key', 14, 0), ('product-final-copy', 12, 12), ('product-final-materials', 10, 22)):
        el = K.find(sl, nm)
        if el is None:
            continue
        gaps.append(gap if els else 0)
        K.set_size(el, sz); K.left_align(el)
        K.add_en(sl, el, EN[(i, nm)], width=260)
        els.append(el)
    K.text_block(sl, els, 640, 284, gaps)

# ------------------------------------------------------------------ 4. slides
K.build_cover(S[1], '有节 · 竹院', 'Nodes · Bamboo Court', '茶家具调研与概念设计', 'Tea Furniture Research & Concept',
              'October 2026', f'{A}/cover_family.png')
K.build_toc(S[2], [('前期调研', 'Preliminary Research'), ('家具与茶室分析', 'Furniture and Tea Spaces'),
                   ('定位与设计', 'Positioning and Design'), ('效果呈现', 'Design Presentation'), ('设计总结', 'Design Summary')],
            art_png=f'{A}/toc_hands.png')
K.build_divider(S[3], '前期调研', 'Preliminary Research', f'{A}/line_actions.png', (200, 282, 560, 150))
K.build_divider(S[11], '家具与茶室分析', 'Furniture and Tea Spaces', f'{A}/chairs_gray.png', (300, 262, 360, 190))
K.build_divider(S[18], '定位与设计', 'Positioning and Design', f'{A}/sketch_table.png', (250, 264, 460, 190),
                tagline='竹有节，器有节，茶有礼节。', tagline_en=TAGLINE_EN)

for i in [4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 17, 19, 20, 21, 22, 23, 24, 25, 26, 40, 41, 42]:
    sl = S[i]
    base(sl, i)
    cn, en = take_header(sl); cr = K.take_text(sl, 'v20-source')
    K.header(sl, cn, en)
    if cr:
        K.credit(sl, cr)
    K.swap_full_renders(sl, CUT)
    if i in CUSTOM:
        CUSTOM[i](sl)
    else:
        K.apply_en(sl, {nm: spec for (si, nm), spec in EN.items() if si == i}, rows=i in ROW_FLOW, tag=f'S{i}')
        for a, b, mode in ANCHORS.get(i, []):
            K.anchor(sl, a, b, mode)
        K.vfit(sl)
    if K.content_bottom(sl) > 494:
        K.log.append(f'S{i}: content bottom {K.content_bottom(sl):.0f}pt')

for el in K.top_shapes(S[4]):                       # archival grayscale for the timeline figures
    sha, _ = K.pic_sha(S[4], el)
    if sha and os.path.exists(f'{A}/gray_{sha}.png'):
        K.replace_blob(S[4], el, f'{A}/gray_{sha}.png')

sl = S[19]                                          # F: hairlines between source / idea / product columns
xs = sorted(K.bbox(el)[0] for el in K.top_shapes(sl) if K.name_of(el).startswith('inspiration-name'))
pnt = K.bbox(K.find(sl, 'inspiration-painting'))
prod_x = min(K.bbox(el)[0] for el in K.top_shapes(sl) if K.name_of(el).startswith('inspiration-product'))
K.line(sl, (pnt[0] + pnt[2] + xs[0]) / 2, pnt[1], (pnt[0] + pnt[2] + xs[0]) / 2, pnt[1] + pnt[3], K.HAIR)
K.line(sl, prod_x - 28, pnt[1], prod_x - 28, pnt[1] + pnt[3], K.HAIR)

sl = S[16]                                          # E: case page
base(sl, 16); take_header(sl); cr = K.take_text(sl, 'v20-source'); K.take_text(sl, 'fuzhou-observation')
K.corner_label(sl, '案例分享')
main, inte, wat, ext = (K.find(sl, n) for n in ('fuzhou-main', 'fuzhou-interior', 'fuzhou-water', 'fuzhou-exterior'))
K.place(main, 60, 104, 400, 225)
cw = (400 - 16) / 3
for k, el in enumerate((inte, wat, ext)):
    K.set_crop_aspect(sl, el, cw / 92); K.place(el, 60 + k * (cw + 8), 337, cw, 92)
case = [K.P1('THE RELIC SHELTER', 14, K.TITLE, spc=1), K.P1('Neri&Hu｜福州茶室', 12, K.BODY, before=6),
        K.P1('木构、院落与界面', 12, K.TITLE, before=22), K.P1(CASE16['lead_en'], 10, K.BODY, ln=14, before=3)]
for k, (cn_line, en_line) in enumerate(CASE16['pairs']):
    case += [K.P1(cn_line, 11, K.BODY, ln=17, before=18 if k == 0 else 10), K.P1(en_line, 9, K.BODY, ln=13, before=3)]
K.tb(sl, 540, 104, 340, 325, case, align='c', anchor='ctr', name='case-text')
K.credit(sl, cr)

for i in (27, 29, 31, 33, 35, 37):
    sketch_page(S[i], i)
for i in (28, 30, 32, 34, 36, 38):
    render_page(S[i], i)

sl = S[39]                                          # scene -> G-style title page
base(sl, 39); take_header(sl); K.take_text(sl, 'scene-context')
hero = K.find(sl, 'scene-hero')
K.build_g(sl, 'Waterside · Tea and Rest', '临水共饮　有顶廊下 · 五人茶席')
K.place(hero, *K.fit(K.pic_aspect(sl, hero), 120, 150, 720, 336))

c04 = K.new_slide(prs); K.build_divider(c04, '效果呈现', 'Design Presentation', f'{A}/lineup.png', (190, 300, 580, 120))
c05 = K.new_slide(prs); K.build_divider(c05, '设计总结', 'Design Summary', f'{A}/line_journey.png', (210, 282, 540, 150))
PRODUCTS = [('Tea Table', '茶台', 'table', (150, 158, 660, 316)), ('Host Chair', '主位椅', 'host', (160, 152, 640, 330)),
            ('Guest Chair', '客位椅', 'guest', (160, 152, 640, 330)), ('Twin Seat', '并坐连坐', 'twin', (150, 158, 660, 316)),
            ('Display Shelf', '展示架', 'shelf', (160, 150, 640, 336)), ('Folded Light', '折光吊灯', 'lamp', (160, 146, 640, 344))]
gpages = []
for en, cn, key, box in PRODUCTS:
    g = K.new_slide(prs); K.build_g(g, en, cn, f'{A}/blur_{key}.png', box); gpages.append(g)
back = K.new_slide(prs)
K.build_cover(back, '有节 · 竹院', 'Nodes · Bamboo Court', '茶家具调研与概念设计', 'Tea Furniture Research & Concept',
              'October 2026', f'{A}/cover_family.png')

order = [S[1], S[2], S[3]] + [S[i] for i in range(4, 11)] + [S[11]] + [S[i] for i in range(12, 18)] + [S[18]] + \
        [S[i] for i in range(19, 27)] + [c04]
for k, i in enumerate((27, 29, 31, 33, 35, 37)):
    order += [gpages[k], S[i], S[i + 1]]
order += [S[39], S[40], c05, S[41], S[42], back]
K.reorder(prs, order)
prs.save(OUT)
K.set_theme_fonts(OUT)
print('\n'.join(sorted(set(K.log))) or 'no warnings')
print('saved', OUT, len(order), 'slides')
