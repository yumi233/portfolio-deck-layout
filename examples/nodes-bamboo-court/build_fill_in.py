"""Fill-in version (待补图版): the new-render build plus a grey slot with an AI prompt for every image that is still missing.

  python -I build_fill_in.py SOURCE.pptx OUT.pptx WORKDIR NEW_IMAGES_DIR
"""
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
SRC, OUT, WORK, NEWIMG = sys.argv[1:5]
SKILL = sys.argv[5] if len(sys.argv) > 5 else os.path.join(HERE, '..', '..')
sys.path.insert(0, os.path.join(SKILL, 'scripts'))
sys.path.insert(0, os.path.join(SKILL, 'examples', 'nodes-bamboo-court'))
sys.path.append(HERE)                      # last: this folder also holds older one-off scripts
import numpy as np
from PIL import Image
from pptx import Presentation
import pptx_kit as K
import assets as AS
import fonts as FN
from en_text import EN, ROW_FLOW, FORM_STUDY_EN, CASE16
from regions import REGIONS

MEDIA, A = os.path.join(WORK, 'media'), os.path.join(WORK, 'assets')
os.makedirs(MEDIA, exist_ok=True); os.makedirs(A, exist_ok=True)
cn_font, _ = FN.find(FN.CN); en_font, en_file = FN.find(FN.EN)
K.configure(en_font=en_font, cn_font=cn_font, en_font_file=en_file)

prs = Presentation(SRC)
prs.slide_width, prs.slide_height = K.E(960), K.E(540)
S = {i + 1: s for i, s in enumerate(prs.slides)}

# ------------------------------------------------------------------ names
THEME = ('节序', 'Nodes in Order')
TAGLINE = ('竹有节，茶有序。', 'Bamboo grows in nodes; tea moves in order.')
P = {   # key: (CN name · category, EN header, EN title for the G page)
    'table': ('围炉 · 茶台', 'Hearthside · Tea Table', 'Hearthside'),
    'host': ('东道 · 主位椅', 'East Seat · Host Chair', 'East Seat'),
    'guest': ('西宾 · 客位椅', 'West Seat · Guest Chair', 'West Seat'),
    'twin': ('比肩 · 并坐连坐', 'Shoulder to Shoulder · Twin Seat', 'Shoulder to Shoulder'),
    'shelf': ('藏器 · 展示架', 'Vessel Keeper · Display Shelf', 'Vessel Keeper'),
    'lamp': ('秉烛 · 折光吊灯', 'Candlelight · Folded Pendant', 'Candlelight'),
}
ORDER = ['table', 'host', 'guest', 'twin', 'shelf', 'lamp']
SKETCH = dict(zip((27, 29, 31, 33, 35, 37), ORDER))
RENDER = {i + 1: k for i, k in SKETCH.items()}

# ------------------------------------------------------------------ 1. images
def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        elif s.shape_type == 13:
            yield s
for sl in prs.slides:
    for s in walk(sl.shapes):
        p = os.path.join(MEDIA, s.image.sha1[:8] + '.bin')
        if not os.path.exists(p):
            open(p, 'wb').write(s.image.blob)
img = lambda h: Image.open(os.path.join(MEDIA, h + '.bin'))
new = lambda k: os.path.join(NEWIMG, f'{k}.jpg')

for k in range(1, 17):                                  # user's new renders: grey studio -> transparent
    if not os.path.exists(f'{A}/new_{k}.png'):
        AS.studio_cutout(Image.open(new(k))).save(f'{A}/new_{k}.png')
cut = lambda k: f'{A}/new_{k}.png'
HERO = {'table': 7, 'host': 12, 'guest': 16, 'twin': 9, 'shelf': 3, 'lamp': 1}
TEASER = {'table': 5, 'host': 12, 'guest': 16, 'twin': 9, 'shelf': 3, 'lamp': 1}
VIEWS = {'table': [(5, '正面', 'Front'), (6, '背面', 'Back')],
         'host': [(13, '正面', 'Front'), (14, '侧面', 'Side'), (15, '背面', 'Back')],
         'guest': [(11, '侧面', 'Side'), (10, '背面', 'Back')],
         'twin': [(8, '正视 · 侧视 · 俯视', 'Front · Side · Top')],
         'shelf': [(4, '正面', 'Front'), (2, '侧面', 'Side')]}
# lamp without most of its cable: the long cable would shrink the lamp body on every page
lamp = Image.open(cut(1)).convert('RGBA'); al = np.asarray(lamp)[..., 3]
widths = (al > 128).sum(1); body_top = int(np.argmax(widths > 0.4 * widths.max()))
body_h = lamp.height - body_top
AS.trim(lamp.crop((0, max(0, body_top - int(0.12 * body_h)), lamp.width, lamp.height))).save(f'{A}/new_lamp_short.png')
hero_path = lambda key: f'{A}/new_lamp_short.png' if key == 'lamp' else cut(HERO[key])
for key, k in TEASER.items():
    src_png = f'{A}/new_lamp_short.png' if key == 'lamp' else cut(k)
    AS.motion_blur(Image.open(src_png)).save(f'{A}/blur_new_{key}.png')
for k, name in ((5, 'table_side'), (12, 'host'), (16, 'guest'), (9, 'twin'), (3, 'shelf')):
    Image.open(cut(k)).save(f'{A}/fam_{name}.png')
Image.open(f'{A}/new_lamp_short.png').save(f'{A}/fam_lamp.png')
for name, spec in {
    'cover_family_v25.png': {'canvas': [480, 320], 'items': [
        {'src': 'fam_shelf.png', 'x': 8, 'y': 6, 'h': 170}, {'src': 'fam_lamp.png', 'x': 420, 'y': 0, 'h': 150},
        {'src': 'fam_table_side.png', 'x': 136, 'y': 52, 'w': 250}, {'src': 'fam_host.png', 'x': 22, 'y': 182, 'h': 118},
        {'src': 'fam_twin.png', 'x': 150, 'y': 203, 'w': 196}, {'src': 'fam_guest.png', 'x': 372, 'y': 186, 'h': 112}]},
    'lineup_v25.png': {'canvas': [560, 112], 'items': [
        {'src': 'fam_shelf.png', 'x': 0, 'y': 0, 'h': 110}, {'src': 'fam_table_side.png', 'x': 88, 'y': 43, 'w': 150},
        {'src': 'fam_host.png', 'x': 250, 'y': 40, 'h': 70}, {'src': 'fam_guest.png', 'x': 320, 'y': 42, 'h': 68},
        {'src': 'fam_twin.png', 'x': 392, 'y': 53, 'w': 118}, {'src': 'fam_lamp.png', 'x': 522, 'y': 0, 'h': 110}]},
}.items():
    AS.compose(spec, A).save(f'{A}/{name}')
# assets kept from the original deck
for h in ['603d77fc', '8d7f3f5b', 'ac8d0326', '69ec03cf', '9a157f22', 'efe0c995']:
    AS.archival_gray(img(h)).save(f'{A}/gray_{h}.png')
AS.trim(img('26b1a91d').convert('RGBA'), thr=10, pad=6).save(f'{A}/line_actions.png')
hands = AS.trim(img('3aa9f4ef').convert('RGBA'), thr=10, pad=6)
AS.lineart(hands.crop((int(hands.width * 0.5), 0, hands.width, hands.height)), gray=138).save(f'{A}/toc_hands.png')
AS.trim(img('1c5f7226').convert('RGBA'), thr=40, pad=4).save(f'{A}/sketch_table.png')
AS.trim(img('91231817').convert('RGBA'), thr=40, pad=4).save(f'{A}/line_journey.png')
c1, c2 = AS.knockout(img('608752ff')), AS.knockout(img('d00cffb0'))
c1 = c1.resize((round(c1.width * 600 / c1.height), 600), Image.LANCZOS)
c2 = c2.resize((round(c2.width * 516 / c2.height), 516), Image.LANCZOS)
can = Image.new('RGBA', (c1.width + 150 + c2.width, 600), (255, 255, 255, 0))
can.alpha_composite(c1, (0, 0)); can.alpha_composite(c2, (c1.width + 150, 600 - c2.height)); can.save(f'{A}/chairs_gray.png')

OLD2NEW = {'5daefaf6': cut(12), '691ab745': cut(16), 'e3b9ab7d': cut(7), 'e95188a2': cut(9),
           '3d1eb110': cut(3), '0f72668c': f'{A}/new_lamp_short.png'}

# ------------------------------------------------------------------ 2. copy changes (the furniture changed)
CN = {
    (14, 'matrix-label-3'): ['节序'],
    (17, 'strategy-copy-2'): ['立柜、杆架与节点，', '各有分工。'],
    (20, 'seat-copy-1'): ['圈椅的围合关系，', '转为弯管靠框与V形扶手。'],
    (23, 'material-copy-0'): ['茶台、扶手与中央几，', '保留木纹与接触温度。'],
    (23, 'material-copy-2'): ['皮革靠背与棉编带座面，', '区分身体接触层。'],
    (24, 'process-copy-2'): ['开料、铣槽、收边与表面处理，', '形成桌面、扶手与小几。'],
    (24, 'process-copy-3'): ['皮革缝制与编带交织，', '接触层与杆架疏密相对。'],
    (26, 'collection-table-label'): ['01 围炉 · 茶台'], (26, 'collection-host-label'): ['02 东道 · 主位椅'],
    (26, 'collection-guest-label'): ['03 西宾 · 客位椅'], (26, 'collection-twin-label'): ['04 比肩 · 并坐连坐'],
    (26, 'collection-shelf-label'): ['05 藏器 · 展示架'], (26, 'collection-light-label'): ['06 秉烛 · 折光吊灯'],
    (27, 'product-form-tag'): ['一柜一杆'],
    (27, 'product-form-copy'): ['立柜与杆架分立两端，', '细杆框架显露连接。', '水槽、茶盘与烧水位嵌入台面。'],
    (28, 'product-final-key'): ['一柜一杆'],
    (28, 'product-final-materials'): ['胡桃木台面 · 不锈钢立柜与杆架 · 嵌入水槽、茶盘与烧水位'],
    (29, 'product-form-copy'): ['官帽椅的上扬搭脑，', '转为胡桃棕皮革靠背。', '胡桃木扶手落在金属套环上。'],
    (30, 'product-final-copy'): ['胡桃棕皮革靠背，胡桃木扶手与横撑。', '外露杆架与转角节点保持清楚轮廓。'],
    (30, 'product-final-materials'): ['胡桃棕皮革靠背', '胡桃木扶手与横撑', '棉编带座面'],
    (31, 'product-form-copy'): ['圈椅的围合关系，', '转为弯管靠框与V形扶手。', '宽编带座面保留疏密。'],
    (32, 'product-final-copy'): ['弯管靠框托住皮革靠背，', 'V形撑托起扶手，侧面保持留隙。'],
    (32, 'product-final-materials'): ['胡桃棕皮革靠背', 'V形扶手撑', '宽编带座面'],
    (33, 'product-form-copy'): ['两只皮革编带座椅与中央几，', '沿同一横杆串接。', '几面放杯，几下收整。'],
    (34, 'product-final-materials'): ['皮革靠背与编带座 · 胡桃木中几 · 共享横杆'],
    (36, 'product-final-materials'): ['两端织面', '四层长虹玻璃', '开放中部'],
    (37, 'product-form-copy'): ['浅色折面交替转向，', '沿四角方杆逐层排列。', '暖光透过薄面与间隙。'],
    (38, 'product-final-materials'): ['纤维折面', '不锈钢方杆', '井字横框'],
}
EN.update({
    (17, 'strategy-copy-2'): ('block', 'Pedestal, rod frames and joints each do their own job.'),
    (20, 'seat-copy-1'): ('block', 'The embrace of the round-back chair becomes bent-tube back loops and V-braced arms.'),
    (23, 'material-copy-0'): ('block', 'Tea table, armrests and centre table keep the grain and warmth of wood.'),
    (23, 'material-copy-2'): ('block', 'Leather backs and cotton-webbing seats mark where the body rests.'),
    (24, 'process-copy-2'): ('block', 'Cutting, routing, edging and finishing make the tabletop, armrests and side table.'),
    (24, 'process-copy-3'): ('block', 'Stitched leather and woven webbing; soft contact layers set against the open frame.'),
    (27, 'product-form-tag'): ('block', 'One Pedestal, One Frame'),
    (27, 'product-form-copy'): ('block', 'A steel pedestal and a rod frame stand at either end; the slender frame shows its joints. Sink, tray and kettle plate sit flush in the top.'),
    (28, 'product-final-key'): ('block', 'One Pedestal, One Frame'),
    (28, 'product-final-materials'): ('block', 'Walnut top · Steel pedestal and rod frame · Inset sink, tray and kettle plate'),
    (29, 'product-form-copy'): ('block', "The upswept crest rail of the official's hat chair becomes a walnut-brown leather back; walnut armrests rest on metal collars."),
    (30, 'product-final-copy'): ('block', 'A walnut-brown leather back with walnut armrests and stretchers. The exposed rod frame and corner joints keep a crisp outline.'),
    (30, 'product-final-materials'): ('block', 'Walnut-brown leather back\nWalnut armrests and stretchers\nCotton-webbing seat'),
    (31, 'product-form-copy'): ('block', 'The embrace of the round-back chair becomes bent-tube back loops and V-braced arms. A wide-webbing seat keeps an open rhythm.'),
    (32, 'product-final-copy'): ('block', 'Bent-tube loops hold the leather back; V-braces lift the walnut arms and leave the sides open.'),
    (32, 'product-final-materials'): ('block', 'Walnut-brown leather back\nV-braced walnut arms\nWide-webbing seat'),
    (33, 'product-form-copy'): ('block', 'Two leather-and-webbing seats and a centre table are strung along one cross rail. Cups go on top; things are tidied away below.'),
    (34, 'product-final-materials'): ('block', 'Leather backs and webbing seats · Walnut centre table · Shared cross rail'),
    (36, 'product-final-materials'): ('block', 'Woven end panels\nFour fluted-glass shelves\nOpen centre'),
    (37, 'product-form-copy'): ('block', 'Pale folded bands turn in alternation, tier by tier along four square rods. Warm light slips through shades and gaps.'),
    (38, 'product-final-materials'): ('block', 'Folded fibre bands\nStainless square rods\nCross-bar tiers'),
})
HEADER = {26: ('节序 · 系列', 'The Collection'), 41: ('竹之节，茶之序', 'Design Summary')}

# ------------------------------------------------------------------ 3. colour language (unchanged from v24)
def chapter(i):
    return 0 if i <= 2 else 1 if i <= 10 else 2 if i <= 17 else 3 if i <= 26 else 4 if i <= 40 else 5
RESEARCH = (1, 2)
TEXT = {'20342F': K.TITLE, '182C27': K.TITLE, '1B302A': K.TITLE, '0D1C1A': K.TITLE,
        '4A514D': K.BODY, '515C55': K.BODY, '7A817B': K.AUX, '727171': K.AUX}
LINE = {'ADB7AD': K.HAIR, 'C8CBC5': K.HAIR, 'C7D0C6': K.HAIR, 'C8D1C8': K.HAIR, 'AFB4AE': K.OUTL, '96AAA7': K.OUTL,
        '20342F': K.MIDL, '0D1C1A': K.MIDL, '1B302A': K.MIDL, '698A77': K.MIDL, 'BCC8BE': 'FFFFFF'}
FILL = {'EDF0E9': 'FFFFFF', 'EEF0E9': 'FFFFFF', 'E3EAE2': 'FFFFFF', 'F7F7F4': 'FFFFFF', 'F6F6F2': 'FFFFFF',
        'DFE4DD': 'E6E6E4', 'E0E5DE': 'E6E6E4', 'BCC8BE': 'C9CACA', 'C7D0C6': 'DCDCDA'}
ACCENTS = {'718779', '8A715A', '82654E', '73929C', '698A77', '604C3F', 'A17B59', '638A99', '628F7A'}
OVERRIDES = {
    4: {'text': {a: K.BODY for a in ('718779', '8A715A', '82654E', '73929C')}},
    5: {'fill': {'628F7A': K.RED, 'A17B59': 'A9686A', '638A99': 'C49A9C'},
        'line': {'628F7A': K.RED, 'A17B59': 'A9686A', '638A99': 'C49A9C'}},
    6: {'fill': {'A9BEA1': 'C79A9C', 'A5C1BA': 'DDBFC0', 'C9BA99': 'E8D6D7', 'B8BDA6': 'B98284', 'B49B57': K.RED,
                 'A96E50': K.RED, '668C8E': K.RED, '678A57': K.RED, '6E7152': K.RED, '8AA8A0': K.RED},
        'line': {'A9BEA1': 'FFFFFF', 'A5C1BA': 'FFFFFF', 'C9BA99': 'FFFFFF', 'B8BDA6': 'FFFFFF', 'C7D0C6': K.HAIR,
                 'B49B57': K.RED, 'A96E50': K.RED}},
    14: {'fill': {'8A715A': K.OUTL, '718779': K.OUTL, '73929C': K.OUTL, '20342F': K.RED},
         'text': {'8A715A': K.AUX, '718779': K.AUX, '73929C': K.AUX}},
}
def base(sl, i):
    K.drop_chrome(sl, names=('Image 2', 'background-softener', 'v20-running', 'v20-page'), texts=('NODES / BAMBOO COURT',))
    K.scale_geometry(sl, 0.8)
    ov = dict(OVERRIDES.get(i, {}))
    ov.setdefault('fill', {})
    for c in ('20342F', '0D1C1A', '1B302A'):
        ov['fill'].setdefault(c, K.RED if chapter(i) in RESEARCH else K.AUX)
    K.recolor(sl, chapter(i) in RESEARCH, TEXT, LINE, FILL, ACCENTS, ov, tag=f'S{i}')
    K.retype(sl)
    K.restyle_tables(sl)
    K.white_bg(sl)
    for (si, nm), lines in CN.items():
        if si == i:
            el = K.find(sl, nm)
            if el is None:
                K.log.append(f'S{i}: CN target {nm} missing')
            else:
                K.set_lines(el, lines)

def take_header(sl, i):
    cn = K.take_text(sl, 'v20-title'); en = K.en_case(K.take_text(sl, 'v20-subtitle'))
    return HEADER.get(i, (cn, en))

def apply_regions(sl, i):
    for (si, nm), (k, cx, cy, w, _) in REGIONS.items():
        if si == i:
            K.set_region(sl, K.find(sl, nm), new(k), cx, cy, w)

# ------------------------------------------------------------------ 4. per-slide layouts (as v24)
def layout_timeline(sl):
    tx, tw = [60, 210, 360, 510, 660, 807], [135, 140, 135, 140, 135, 93]
    lead = K.find(sl, 'history-lead')
    K.add_en(sl, lead, EN[(4, 'history-lead')], width=900 - K.bbox(lead)[0]); K.move_to(lead, y=72)
    def stack(k, y):
        for nm, gap in (('date', 6), ('name', 8), ('copy', 8), ('meta', 0)):
            el = K.find(sl, f'history-{nm}-{k}')
            if nm == 'copy':
                K.add_en(sl, el, EN[(4, f'history-copy-{k}')], width=tw[k])
            else:
                K.fit_height(el)
            K.move_to(el, tx[k], y); y += K.bbox(el)[3] + gap
        return y
    imgs = {k: K.find(sl, f'history-image-{k}') for k in range(6)}
    top = K.bbox(lead)[1] + K.bbox(lead)[3] + 12
    row1 = max(stack(k, top) for k in (0, 2, 4))
    spine_y = max(row1, top + max(K.bbox(imgs[k])[3] for k in (1, 3, 5))) + 14
    for k in (1, 3, 5):
        K.move_to(imgs[k], y=spine_y - 14 - K.bbox(imgs[k])[3])
    sp = K.find(sl, 'history-spine'); K.move_to(sp, y=spine_y)
    for k in range(6):
        t = K.find(sl, f'history-tick-{k}'); K.move_to(t, y=spine_y + K.bbox(sp)[3] / 2 - K.bbox(t)[3] / 2)
    row2 = max(stack(k, spine_y + 16) for k in (1, 3, 5))
    for k in (0, 2, 4):
        K.move_to(imgs[k], y=spine_y + 14)
    end = max(row2, max(K.bbox(imgs[k])[1] + K.bbox(imgs[k])[3] for k in (0, 2, 4))) + 14
    K.move_to(K.find(sl, 'history-end-rule'), y=end)
    focus = K.find(sl, 'history-design-focus')
    K.add_en(sl, focus, EN[(4, 'history-design-focus')], width=500); K.move_to(focus, y=end + 10)

def layout_process(sl):
    rows = []
    for k in range(5):
        cp = K.find(sl, f'process-copy-{k}')
        K.add_en(sl, cp, EN[(24, f'process-copy-{k}')], width=232)
        rows.append(max(K.bbox(cp)[3], K.bbox(K.find(sl, f'process-image-{k}'))[3]))
    y = 74 + max(0, (416 - sum(rows) - 14 * 4) / 2)
    for k, rh in enumerate(rows):
        for nm in ('image', 'no', 'name', 'rule', 'copy'):
            el = K.find(sl, f'process-{nm}-{k}'); K.move_to(el, y=y + rh / 2 - K.bbox(el)[3] / 2)
        y += rh + 14

def swap_collection(sl):
    for nm, path in (('collection-table', cut(7)), ('collection-host', cut(12)), ('collection-guest', cut(16)),
                     ('collection-twin', cut(9)), ('collection-shelf', cut(3)), ('collection-light', f'{A}/new_lamp_short.png')):
        K.swap_cutout(sl, K.find(sl, nm), path)

CUSTOM = {4: layout_timeline, 24: layout_process}
PRE = {26: swap_collection}
ANCHORS = {23: [(f'material-no-{k}', f'material-name-{k}', 1) for k in range(4)],
           12: [(f'furniture-point-{k}', 'furniture-spine', 'c') for k in range(3)],
           21: [('user-context-title', 'user-context-image', ('below', 18)), ('user-context-copy', 'user-context-title', ('below', 8))]}

def sketch_page(sl, i):
    base(sl, i)
    K.take_text(sl, 'product-form-name'); K.take_text(sl, 'product-form-en')
    cn, en, _ = P[SKETCH[i]]
    K.header(sl, cn, en)
    K.place(K.find(sl, 'product-form-plate'), 56, 92, 552, 368)
    kick, tag, cp = (K.find(sl, n) for n in ('product-form-kicker', 'product-form-tag', 'product-form-copy'))
    for el, sz in ((kick, 10), (tag, 14), (cp, 12)):
        K.left_align(el); K.set_size(el, sz)
    K.add_en(sl, kick, ('inline', FORM_STUDY_EN))
    K.add_en(sl, tag, EN[(i, 'product-form-tag')], width=260)
    K.add_en(sl, cp, EN[(i, 'product-form-copy')], width=260)
    K.text_block(sl, [kick, tag, cp], 650, 276, [0, 8, 16])

def render_page(sl, i):
    base(sl, i)
    take_header(sl, i)
    key = RENDER[i]; cn, en, _ = P[key]
    K.header(sl, cn, en)
    K.swap_cutout(sl, K.find(sl, 'product-final'), hero_path(key), box=(64, 92, 496, 384))
    els, gaps = [], []
    for nm, sz, gap in (('product-final-key', 14, 0), ('product-final-copy', 12, 12), ('product-final-materials', 10, 22)):
        el = K.find(sl, nm)
        if el is None:
            continue
        gaps.append(gap if els else 0)
        K.set_size(el, sz); K.left_align(el)
        K.add_en(sl, el, EN[(i, nm)], width=260)
        els.append(el)
    K.text_block(sl, els, 640, 284, gaps)

# ------------------------------------------------------------------ 5. slides
K.build_cover(S[1], *THEME, '茶家具调研与概念设计', 'Tea Furniture Research & Concept', 'October 2026', f'{A}/cover_family_v25.png')
K.build_toc(S[2], [('前期调研', 'Preliminary Research'), ('家具与茶室分析', 'Furniture and Tea Spaces'),
                   ('定位与设计', 'Positioning and Design'), ('效果呈现', 'Design Presentation'), ('设计总结', 'Design Summary')],
            art_png=f'{A}/toc_hands.png')
K.build_divider(S[3], '前期调研', 'Preliminary Research', f'{A}/line_actions.png', (200, 282, 560, 150))
K.build_divider(S[11], '家具与茶室分析', 'Furniture and Tea Spaces', f'{A}/chairs_gray.png', (300, 262, 360, 190))
K.build_divider(S[18], '定位与设计', 'Positioning and Design', f'{A}/sketch_table.png', (250, 264, 460, 190),
                tagline=TAGLINE[0], tagline_en=TAGLINE[1])

for i in [4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 17, 19, 20, 21, 22, 23, 24, 25, 26, 40, 41, 42]:
    sl = S[i]
    base(sl, i)
    cn, en = take_header(sl, i); cr = K.take_text(sl, 'v20-source')
    K.header(sl, cn, en)
    if cr:
        K.credit(sl, cr)
    if i in PRE:
        PRE[i](sl)
    K.swap_full_renders(sl, OLD2NEW)
    apply_regions(sl, i)
    if i in CUSTOM:
        CUSTOM[i](sl)
    else:
        K.apply_en(sl, {nm: spec for (si, nm), spec in EN.items() if si == i}, rows=i in ROW_FLOW, tag=f'S{i}')
        for a, b, mode in ANCHORS.get(i, []):
            K.anchor(sl, a, b, mode)
        K.vfit(sl)
    if K.content_bottom(sl) > 494:
        K.log.append(f'S{i}: content bottom {K.content_bottom(sl):.0f}pt')

for el in K.top_shapes(S[4]):
    sha, _ = K.pic_sha(S[4], el)
    if sha and os.path.exists(f'{A}/gray_{sha}.png'):
        K.replace_blob(S[4], el, f'{A}/gray_{sha}.png')

sl = S[19]
xs = sorted(K.bbox(el)[0] for el in K.top_shapes(sl) if K.name_of(el).startswith('inspiration-name'))
pnt = K.bbox(K.find(sl, 'inspiration-painting'))
prod_x = min(K.bbox(el)[0] for el in K.top_shapes(sl) if K.name_of(el).startswith('inspiration-product'))
K.line(sl, (pnt[0] + pnt[2] + xs[0]) / 2, pnt[1], (pnt[0] + pnt[2] + xs[0]) / 2, pnt[1] + pnt[3], K.HAIR)
K.line(sl, prod_x - 28, pnt[1], prod_x - 28, pnt[1] + pnt[3], K.HAIR)

sl = S[16]
base(sl, 16); take_header(sl, 16); cr = K.take_text(sl, 'v20-source'); K.take_text(sl, 'fuzhou-observation')
K.corner_label(sl, '案例分享')
main, inte, wat, ext = (K.find(sl, n) for n in ('fuzhou-main', 'fuzhou-interior', 'fuzhou-water', 'fuzhou-exterior'))
K.place(main, 60, 104, 400, 225)
cw = (400 - 16) / 3
for k, el in enumerate((inte, wat, ext)):
    K.set_crop_aspect(sl, el, cw / 92); K.place(el, 60 + k * (cw + 8), 337, cw, 92)
case = [K.P1('THE RELIC SHELTER', 14, K.TITLE, spc=1), K.P1('Neri&Hu｜福州茶室', 12, K.BODY, before=6),
        K.P1('木构、院落与界面', 12, K.TITLE, before=22), K.P1(CASE16['lead_en'], 10, K.BODY, ln=14, before=3)]
for k, (cn_line, en_line) in enumerate(CASE16['pairs']):
    case += [K.P1(cn_line, 11, K.BODY, ln=17, before=18 if k == 0 else 10), K.P1(en_line, 9, K.BODY, ln=13, before=3)]
K.tb(sl, 540, 104, 340, 325, case, align='c', anchor='ctr', name='case-text')
K.credit(sl, cr)

for i in SKETCH:
    sketch_page(S[i], i)
for i in RENDER:
    render_page(S[i], i)

sl = S[39]
base(sl, 39); take_header(sl, 39); K.take_text(sl, 'scene-context')
hero = K.find(sl, 'scene-hero')
K.build_g(sl, 'Waterside · Tea and Rest', '临水共饮　有顶廊下 · 五人茶席')
K.place(hero, *K.fit(K.pic_aspect(sl, hero), 120, 150, 720, 336))

# ------------------------------------------------------------------ 6. v26: fill-in slots for every missing image
from slots import (ANCHOR, REF, NUM, views as P_VIEWS, exploded as P_EXPLODE, DETAIL, detail as P_DETAIL,
                       scene as P_SCENE, sketch as P_SKETCH, ANGLE, SERIES, MATERIALS)
SLOTDIR = os.path.join(WORK, 'slots')
_, CN_FILE = FN.find(FN.CN)
def slot(sl, x, y, w, h, sid, head, ref, prompt):
    return K.add_slot(sl, x, y, w, h, sid, head, ref, prompt, SLOTDIR, font_file=CN_FILE)
def label(sl, x, y, w, cn, en, align='c'):
    return K.tb(sl, x, y, w, 26, [K.P1(cn, 10, K.BODY), K.P1(en, 8, K.AUX, before=1)], align=align, name='slot-label')

# 6a. outdated pictures keep their place; a red-bordered card on top says what to generate
for i, key in SKETCH.items():
    if key in ('shelf', 'lamp'):
        continue                                    # their sketches already match the new design
    nn = NUM[key]; x, y, w, h = K.bbox(K.find(S[i], 'product-form-plate'))
    K.add_old_card(S[i], x + 10, y + h - 158, 360, 148, f'{nn}-K', '形态推敲草图 · 3:2 · 2400×1600（旧稿：旧版造型）', REF[key], P_SKETCH(key))
x, y, w, h = K.bbox(K.find(S[39], 'scene-hero'))
K.add_old_card(S[39], x + 12, y + h - 170, 420, 158, '00-M1', SERIES['00-M1'][0] + '（旧场景：旧版藤编家具）', SERIES['00-M1'][1], SERIES['00-M1'][2])
for i, nm in ((21, 'user-context-image'), (41, 'summary-scene')):
    x, y, w, h = K.bbox(K.find(S[i], nm))
    K.add_old_card(S[i], x + 8, y + h - 96, min(w - 16, 300), 88, '00-M1', '旧场景 · 待替换', '',
                   '这张旧图里还是旧版藤编家具。等 00-M1 大场景（第 71 页的提示词）生成好后，用同一张图替换这里。',
                   steps='① 先按第 71 页生成 00-M1　② 删除本提示框　③ 右键下面的旧图 →「更改图片」，再按需裁剪')

# 6b. new pages
def page_dims(key):
    cn, en, _ = P[key]; nn = NUM[key]
    sl = K.new_slide(prs); K.white_bg(sl); K.header(sl, cn, en)
    im = Image.open(hero_path(key)); K.pic(sl, hero_path(key), *K.fit(im.width / im.height, 56, 105, 330, 330), name='dims-render')
    K.tb(sl, 412, 88, 488, 14, [{'runs': [('三视图 · 尺寸', 10, K.AUX), ('  Orthographic Views · mm', 8.5, K.AUX)]}], name='dims-kicker')
    slot(sl, 412, 108, 488, 325, f'{nn}-V', '三视图 + 尺寸 · 3:2 · 2400×1600（最好 3D 导出）', REF[key], P_VIEWS(key))
    K.tb(sl, 412, 445, 488, 30, [K.P1('尺寸　L ______ × W ______ × H ______ mm（请填写真实尺寸）', 10, K.BODY),
                                 K.P1('Dimensions in millimetres', 8, K.AUX, before=1)], name='dims-text')
    return sl

def page_detail(key):
    cn, en, _ = P[key]; nn = NUM[key]
    sl = K.new_slide(prs); K.white_bg(sl); K.header(sl, cn, en)
    slot(sl, 56, 84, 304, 380, f'{nn}-X', '爆炸图 · 4:5 · 1600×2000', REF[key], P_EXPLODE(key))
    label(sl, 56, 470, 304, '结构拆解', 'Exploded View', align='l')
    K.tb(sl, 400, 158, 500, 14, [{'runs': [('节点与材质', 10, K.AUX), ('  Joints and Materials', 8.5, K.AUX)]}], name='detail-kicker')
    for d in range(3):
        x = 400 + d * (154 + 19)
        slot(sl, x, 182, 154, 154, f'{nn}-D{d + 1}', '细节特写 · 1:1 · 1200×1200', REF[key], P_DETAIL(key, d))
        label(sl, x - 10, 344, 174, DETAIL[key][d][1], DETAIL[key][d][2])
    return sl

def page_scene(key):
    nn = NUM[key]
    sl = K.new_slide(prs); K.white_bg(sl)
    slot(sl, 135, 30, 320, 480, f'{nn}-Sa', '场景图（竖）· 2:3 · 1600×2400', REF[key], P_SCENE(key, 0))
    slot(sl, 495, 98, 330, 412.5, f'{nn}-Sb', '场景图 · 4:5 · 1600×2000', REF[key], P_SCENE(key, 1))
    return sl

def page_views_mixed(key, items, max_h=300, gap=48):
    """Views page where some angles are still slots. items: ('png', path, cn, en) | ('slot', sid)."""
    cn, en, _ = P[key]
    sl = K.new_slide(prs); K.white_bg(sl); K.header(sl, cn, en)
    asp = []
    for it in items:
        if it[0] == 'png':
            im = Image.open(it[1]); asp.append(im.width / im.height)
        else:
            asp.append(ANGLE[it[1]][2])
    h = min(max_h, (840 - gap * (len(items) - 1)) / sum(asp))
    x = 480 - (h * sum(asp) + gap * (len(items) - 1)) / 2
    y = 74 + (416 - (h + 36)) / 2
    for it, a in zip(items, asp):
        w = h * a
        if it[0] == 'png':
            K.pic(sl, it[1], x, y, w=w, h=h, name='view-image'); lc, le = it[2], it[3]
        else:
            k2, head, _, tpl, lc, le = ANGLE[it[1]]
            slot(sl, x, y, w, h, it[1], head, REF[k2], tpl.format(a=ANCHOR[k2]))
        label(sl, x - 30, y + h + 10, w + 60, lc, le)
        x += w + gap
    return sl

def page_materials():
    sl = K.new_slide(prs); K.white_bg(sl); K.header(sl, '节序 · 材料', 'The Materials')
    slot(sl, 60, 92, 370, 370, '00-B', SERIES['00-B'][0], SERIES['00-B'][1], SERIES['00-B'][2])
    y = 277 - (6 * 41 + 5 * 14) / 2
    for n, (cn_, en_, use_cn, use_en) in enumerate(MATERIALS, 1):
        K.tb(sl, 480, y, 420, 41, [{'runs': [(f'{n:02d}  ', 9, K.AUX), (cn_, 11, K.TITLE), ('　', 11, K.TITLE), (en_, 9, K.BODY)], 'ln': 15},
                                  {'runs': [(use_cn, 10, K.BODY), ('　', 10, K.BODY), (use_en, 8.5, K.AUX)], 'ln': 15, 'before': 3}],
             name=f'material-{n}')
        y += 41 + 14
    return sl

def page_bigscene(sid):
    sl = K.new_slide(prs); K.white_bg(sl)
    slot(sl, 80, 45, 800, 450, sid, *SERIES[sid])
    return sl

c04 = K.new_slide(prs); K.build_divider(c04, '效果呈现', 'Design Presentation', f'{A}/lineup_v25.png', (190, 300, 580, 120))
c05 = K.new_slide(prs); K.build_divider(c05, '设计总结', 'Design Summary', f'{A}/line_journey.png', (210, 282, 540, 150))
BOX = {'table': (150, 158, 660, 316), 'host': (160, 152, 640, 330), 'guest': (160, 152, 640, 330),
       'twin': (150, 158, 660, 316), 'shelf': (160, 150, 640, 336), 'lamp': (160, 146, 640, 344)}
mat = page_materials()
extra = {}
for key in ORDER:
    cn, en, title = P[key]
    pages = []
    g = K.new_slide(prs); K.build_g(g, title, cn, f'{A}/blur_new_{key}.png', BOX[key])
    if key == 'guest':
        v = page_views_mixed(key, [('slot', '03-A'), ('png', cut(11), '侧面', 'Side'), ('png', cut(10), '背面', 'Back')])
    elif key == 'lamp':
        v = page_views_mixed(key, [('png', f'{A}/new_lamp_short.png', '正面', 'Front'), ('slot', '06-A1'), ('slot', '06-A2')])
    else:
        v = K.new_slide(prs); K.white_bg(v); K.header(v, cn, en)
        K.build_views(v, [(cut(k), a, b) for k, a, b in VIEWS[key]], max_h=300 if key != 'twin' else 380)
    extra[key] = (g, v, page_dims(key), page_detail(key), page_scene(key))
m2, m3 = page_bigscene('00-M2'), page_bigscene('00-M3')
back = K.new_slide(prs)
K.build_cover(back, *THEME, '茶家具调研与概念设计', 'Tea Furniture Research & Concept', 'October 2026', f'{A}/cover_family_v25.png')

order = [S[1], S[2], S[3]] + [S[i] for i in range(4, 11)] + [S[11]] + [S[i] for i in range(12, 18)] + [S[18]] + \
        [S[i] for i in range(19, 27)] + [mat, c04]
for i, key in SKETCH.items():
    g, v, dims, det, sc = extra[key]
    order += [g, S[i], S[i + 1], v, dims, det, sc]
order += [S[39], m2, m3, S[40], c05, S[41], S[42], back]
K.reorder(prs, order)
prs.save(OUT)
K.set_theme_fonts(OUT)

# page map for the checklist
import re as _re
pmap = {}
for n, s in enumerate(Presentation(OUT).slides, 1):
    for el in K.top_shapes(s):
        m = _re.match(r'(?:待替换图片|提示词-替换后删除)-(.+)$', K.name_of(el))
        if m:
            pmap.setdefault(m.group(1), set()).add(n)
import json
json.dump({k: sorted(v) for k, v in pmap.items()}, open(os.path.join(WORK, 'slot_pages.json'), 'w'), ensure_ascii=False, indent=1)
print('slots:', len(pmap))

# no old product render may survive (sketch plates and scenes are kept on purpose)
OLD = set(OLD2NEW) | {'c88c4ce8'}
left = []
for n, s in enumerate(Presentation(OUT).slides, 1):
    for el in K.top_shapes(s):
        sha, _ = K.pic_sha(s, el)
        if sha in OLD:
            left.append(f'slide {n} {K.name_of(el)} {sha}')
print('old renders left:', left or 'none')
print('\n'.join(sorted(set(K.log))) or 'no warnings')
print('saved', OUT, len(order), 'slides')
