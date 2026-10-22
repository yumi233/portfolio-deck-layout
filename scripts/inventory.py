"""Read a deck before planning: page size, every slide's text/pictures/shapes, and where each image is reused.

  python inventory.py deck.pptx [--slides 3-14] [--images-out DIR]

Shape lines:  T x,y wxh size #colour [group path]: text     I x,y wxh (px) crop=(l,t,r,b) sha=xxxxxxxx
Coordinates are points in the SOURCE deck. --images-out saves each distinct image once as <sha>.<ext>
so you can look at them and build cut-outs.
"""
import sys, argparse, os, collections
from pptx import Presentation

P = 12700


def flat(shapes, path=''):
    for s in shapes:
        if s.shape_type == 6:
            yield from flat(s.shapes, path + s.name + '/')
        else:
            yield path, s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx'); ap.add_argument('--slides'); ap.add_argument('--images-out')
    a = ap.parse_args()
    prs = Presentation(a.pptx)
    lo, hi = (map(int, a.slides.split('-')) if a.slides else (1, len(prs.slides)))
    print(f'page {prs.slide_width / P:.0f} x {prs.slide_height / P:.0f} pt, {len(prs.slides)} slides, '
          f'scale to 960 wide = {960 / (prs.slide_width / P):.3f}')
    uses = collections.defaultdict(list); blobs = {}
    for i, sl in enumerate(prs.slides, 1):
        show = lo <= i <= hi
        if show:
            print(f'\n=== S{i}  layout={sl.slide_layout.name}')
        small = 0
        for path, s in flat(sl.shapes):
            g = f'{s.left / P:.0f},{s.top / P:.0f} {s.width / P:.0f}x{s.height / P:.0f}' if s.left is not None else '-'
            if s.shape_type == 13:
                sha = s.image.sha1[:8]; uses[sha].append(i); blobs[sha] = (s.image.blob, s.image.ext, s.image.size)
                if show:
                    print(f' I {g} {s.image.size} crop=({s.crop_left:.2f},{s.crop_top:.2f},{s.crop_right:.2f},{s.crop_bottom:.2f}) '
                          f'sha={sha} [{path}{s.name}]')
            elif s.has_text_frame and s.text_frame.text.strip():
                if show:
                    r = [r for pa in s.text_frame.paragraphs for r in pa.runs]
                    sz = r[0].font.size.pt if r and r[0].font.size else '?'
                    try:
                        c = r[0].font.color.rgb
                    except Exception:
                        c = '-'
                    t = s.text_frame.text.strip().replace('\n', ' ⏎ ').replace('\x0b', ' ⏎ ')
                    print(f' T {g} {sz}pt #{c} [{path}{s.name}]: {t}')
            elif s.shape_type == 5 and s.width and s.width < 20 * P and s.height < 20 * P:
                small += 1
            elif show:
                print(f' S {g} type={s.shape_type} [{path}{s.name}]')
        if show and small:
            print(f'  (+{small} small freeforms — maps/icons drawn as vectors; keep them, do not rebuild)')
        if show and sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            print('  notes:', sl.notes_slide.notes_text_frame.text.strip()[:120].replace('\n', ' / '))
    print('\n=== images (sha: size, slides using it)')
    for sha, sl in sorted(uses.items(), key=lambda kv: -len(set(kv[1]))):
        print(f' {sha} {blobs[sha][2]} {blobs[sha][1]} -> slides {sorted(set(sl))}')
        if a.images_out:
            os.makedirs(a.images_out, exist_ok=True)
            open(os.path.join(a.images_out, f'{sha}.{blobs[sha][1]}'), 'wb').write(blobs[sha][0])


if __name__ == '__main__':
    main()
