"""pptx_kit — helpers for laying out any deck in the minimal oriental-serif system (960 x 540 pt).
Works for every project type (product, space, brand, UI, illustration, fashion, photography, research):
the functions know about pages, type, colour and images, not about what the work is.

Import it from a per-deck build script:

    import sys; sys.path.insert(0, '<skill>/scripts')
    import pptx_kit as K
    K.configure(en_font='Hoefler Text', cn_font='Songti SC', en_font_file='/System/Library/Fonts/Supplemental/Hoefler Text.ttc')
    K.set_accent('6F3F2A', 'D8BFB2')     # from `assets.py accent <main images>`; default is tea red

Every coordinate below is in points on the 960 x 540 page. See references/page-templates.md for where each
template puts things and references/build-recipes.md for the order to call these in.
"""
import io, os, re, math, copy, zipfile, shutil, hashlib
from lxml import etree
from PIL import Image, ImageFont
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

P = 12700                                   # EMU per pt
PAGE_W, PAGE_H = 960, 540

# palette — references/layout-rules.md §色彩
TITLE, BODY, AUX = '231815', '3E3A39', '595757'
HAIR, MIDL, OUTL = 'DCDCDA', '7C7979', 'B5B5B6'
RED, REDL = '8C2C2C', 'CBA7A9'              # the accent pair; default tea red, replaced by set_accent()
TOC_BG = 'FAFAFA'


def set_accent(main, light=None):
    """Use the project's own accent (derived from its main colour) instead of tea red.
    Every helper reads RED / REDL at call time, so call this once before building."""
    global RED, REDL
    RED = main.upper().lstrip('#')
    REDL = (light or REDL).upper().lstrip('#')

CFG = {'en_font': 'Hoefler Text', 'cn_font': 'Songti SC', 'en_font_file': None}
_FONT = {}
log = []                                    # build warnings (unmapped colours, overflow, missing shapes)


def configure(en_font=None, cn_font=None, en_font_file=None):
    if en_font: CFG['en_font'] = en_font
    if cn_font: CFG['cn_font'] = cn_font
    if en_font_file: CFG['en_font_file'] = en_font_file
    _FONT.clear()


def E(v):
    return Emu(int(round(v * P)))

# ----------------------------------------------------------------------------- shape access
def name_of(el):
    c = el.find('.//' + qn('p:cNvPr'))
    return c.get('name') if c is not None else ''

def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('a:t')))

def top_shapes(slide):
    return [el for el in slide.shapes._spTree
            if etree.QName(el).localname in ('sp', 'pic', 'grpSp', 'graphicFrame', 'cxnSp')]

def find(slide, name):
    for el in top_shapes(slide):
        if name_of(el) == name:
            return el
    return None

def remove(el):
    if el is not None:
        el.getparent().remove(el)

def clear(slide):
    for el in top_shapes(slide):
        remove(el)

def top_xfrm(el):
    t = etree.QName(el).localname
    if t == 'grpSp':
        return el.find(qn('p:grpSpPr')).find(qn('a:xfrm'))
    if t == 'graphicFrame':
        return el.find(qn('p:xfrm'))
    return el.find(qn('p:spPr')).find(qn('a:xfrm'))

def bbox(el):
    x = top_xfrm(el)
    o, e = x.find(qn('a:off')), x.find(qn('a:ext'))
    return (int(o.get('x')) / P, int(o.get('y')) / P, int(e.get('cx')) / P, int(e.get('cy')) / P)

def shift(el, dx=0, dy=0):
    o = top_xfrm(el).find(qn('a:off'))
    o.set('x', str(int(o.get('x')) + int(dx * P)))
    o.set('y', str(int(o.get('y')) + int(dy * P)))

def move_to(el, x=None, y=None):
    bx, by, bw, bh = bbox(el)
    shift(el, dx=(x - bx) if x is not None else 0, dy=(y - by) if y is not None else 0)

def place(el, x, y, w, h):
    """Move/resize a shape; single-child wrapper groups are resized together with their child."""
    t = etree.QName(el).localname
    xf = top_xfrm(el)
    for tag in ('a:off', 'a:chOff'):
        n = xf.find(qn(tag))
        if n is not None:
            n.set('x', str(int(x * P))); n.set('y', str(int(y * P)))
    for tag in ('a:ext', 'a:chExt'):
        n = xf.find(qn(tag))
        if n is not None:
            n.set('cx', str(int(w * P))); n.set('cy', str(int(h * P)))
    if t == 'grpSp':
        for ch in el:
            if etree.QName(ch).localname in ('sp', 'pic', 'grpSp'):
                place(ch, x, y, w, h)

def fit(aspect, x, y, w, h):
    """Largest box of `aspect` (w/h) centred in (x, y, w, h)."""
    if w / h > aspect:
        nw = h * aspect
        return x + (w - nw) / 2, y, nw, h
    nh = w / aspect
    return x, y + (h - nh) / 2, w, nh

# ----------------------------------------------------------------------------- pictures
def first_pic(el):
    return el if etree.QName(el).localname == 'pic' else el.find('.//' + qn('p:pic'))

def _blip_part(slide, el):
    p = first_pic(el)
    rid = p.find('.//' + qn('a:blip')).get(qn('r:embed'))
    return p, slide.part.related_part(rid)

def pic_aspect(slide, el):
    """Aspect (w/h) of the visible, cropped image."""
    p, part = _blip_part(slide, el)
    W, H = Image.open(io.BytesIO(part.blob)).size
    sr = p.find('.//' + qn('a:srcRect'))
    l = t = r = b = 0
    if sr is not None:
        l, t, r, b = (int(sr.get(k, 0)) / 100000 for k in ('l', 't', 'r', 'b'))
    return (W * (1 - l - r)) / (H * (1 - t - b))

def pic_sha(slide, el):
    """(sha1[:8] of the image blob, [l,t,r,b] crop) — use it to recognise the same render across slides."""
    if first_pic(el) is None:
        return None, None
    p, part = _blip_part(slide, el)
    sr = p.find('.//' + qn('a:srcRect'))
    crop = [int(sr.get(k, 0)) / 100000 for k in ('l', 't', 'r', 'b')] if sr is not None else [0, 0, 0, 0]
    return hashlib.sha1(part.blob).hexdigest()[:8], crop

def set_crop_aspect(slide, el, target):
    """Re-crop (centred on the current crop) so the visible image has aspect `target`."""
    p, part = _blip_part(slide, el)
    W, H = Image.open(io.BytesIO(part.blob)).size
    bf = p.find(qn('p:blipFill'))
    sr = bf.find(qn('a:srcRect'))
    if sr is None:
        sr = etree.Element(qn('a:srcRect'))
        bf.find(qn('a:blip')).addnext(sr)
    l, t, r, b = (int(sr.get(k, 0)) / 100000 for k in ('l', 't', 'r', 'b'))
    vw, vh = W * (1 - l - r), H * (1 - t - b)
    cx, cy = W * (l + (1 - l - r) / 2), H * (t + (1 - t - b) / 2)
    if vw / vh > target:
        vw = vh * target
    else:
        vh = vw / target
    vals = ((cx - vw / 2) / W, (cy - vh / 2) / H, 1 - (cx + vw / 2) / W, 1 - (cy + vh / 2) / H)
    for k, v in zip(('l', 't', 'r', 'b'), vals):
        sr.set(k, str(int(round(v * 100000))))

def replace_blob(slide, el, path):
    """Point a picture at a new image file (e.g. its grayscale version), keeping geometry and crop."""
    p = first_pic(el)
    _, rid = slide.part.get_or_add_image_part(path)
    p.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)

def pic(slide, path, x, y, w=None, h=None, name=None):
    im = Image.open(path)
    a = im.width / im.height
    if w is None:
        w = h * a
    if h is None:
        h = w / a
    shp = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    if name:
        shp.name = name
    return shp

def swap_cutout(slide, el, cutout_path, box=None, inset=0.88):
    """Replace a white-backdrop render with its transparent cut-out at the same z-order.
    box: target (x,y,w,h); default = the old box shrunk by `inset` (cut-outs have no margin)."""
    x, y, w, h = box or bbox(el)
    if box is None:
        x, y, w, h = x + w * (1 - inset) / 2, y + h * (1 - inset) / 2, w * inset, h * inset
    im = Image.open(cutout_path)
    new = pic(slide, cutout_path, *fit(im.width / im.height, x, y, w, h), name=name_of(el) or 'cutout')
    el.addprevious(new._element)
    remove(el)
    return new._element

def swap_full_renders(slide, renders, max_crop=0.3):
    """renders: {sha8: cutout_path}. Swaps every picture that shows a whole product (small crop)."""
    for el in top_shapes(slide):
        sha, crop = pic_sha(slide, el)
        if sha in renders and crop[0] + crop[2] < max_crop and crop[1] + crop[3] < max_crop:
            swap_cutout(slide, el, renders[sha])

# ----------------------------------------------------------------------------- text
def set_fonts(rPr):
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        for n in rPr.findall(qn(tag)):
            rPr.remove(n)
    succ = [rPr.find(qn(t)) for t in ('a:sym', 'a:hlinkClick', 'a:hlinkMouseOver', 'a:rtl', 'a:extLst')]
    succ = [s for s in succ if s is not None]
    for tag, face in (('a:latin', CFG['en_font']), ('a:ea', CFG['cn_font']), ('a:cs', CFG['cn_font'])):
        n = etree.Element(qn(tag)); n.set('typeface', face)
        if succ:
            succ[0].addprevious(n)
        else:
            rPr.append(n)

def tb(slide, x, y, w, h, paras, align='l', anchor='t', name=None, wrap=True):
    """New text box. paras: [{'runs': [(text, size, color[, spc])], 'ln': pitch_pt, 'before': pt}]."""
    shp = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        shp.name = name
    tf = shp.text_frame
    bp = tf._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    bp.set('wrap', 'square' if wrap else 'none')
    bp.set('anchor', anchor)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if para.get('ln'):
            p.line_spacing = Pt(para['ln'])
        if para.get('before'):
            p.space_before = Pt(para['before'])
        for run in para['runs']:
            text, size, color = run[:3]
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
            rPr = r._r.get_or_add_rPr()
            rPr.set('lang', 'zh-CN')
            if len(run) > 3 and run[3]:
                rPr.set('spc', str(int(run[3] * 100)))
            set_fonts(rPr)
    return shp

def P1(text, size, color, ln=None, before=None, spc=None):
    return {'runs': [(text, size, color, spc)], 'ln': ln, 'before': before}

def line(slide, x1, y1, x2, y2, color, w=0.5, name=None):
    c = slide.shapes.add_connector(1, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(w)
    if name:
        c.name = name
    return c

def white_bg(slide, color='FFFFFF'):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = RGBColor.from_string(color)

def left_align(el):
    for p in el.iter(qn('a:p')):
        pPr = p.find(qn('a:pPr'))
        if pPr is None:
            pPr = etree.Element(qn('a:pPr')); p.insert(0, pPr)
        pPr.set('algn', 'l')

def set_size(el, size):
    for rp in el.iter(qn('a:rPr')):
        rp.set('sz', str(int(round(size * 100))))

def en_case(s):
    """'TEA CULTURE / A HISTORICAL TIMELINE' -> 'Tea Culture · A Historical Timeline'."""
    small = {'a', 'an', 'and', 'of', 'the', 'to', 'in', 'on', 'for', 'at', 'by', 'with', 'from', 'across'}
    segs = []
    for seg in re.split(r'\s*/\s*', s.strip()):
        ws = seg.lower().split()
        segs.append(' '.join(w if (i and w in small) else w[:1].upper() + w[1:] for i, w in enumerate(ws)))
    return ' · '.join(segs)

# fixed furniture of content pages
def header(slide, cn, en):
    """H2 page title, top-left: CN 14pt + two ideographic spaces + EN 12pt on one line."""
    runs = [(cn, 14, TITLE)]
    if en:
        runs += [('　　', 12, TITLE), (en, 12, BODY)]
    return tb(slide, 33, 25, 700, 20, [{'runs': runs}], name='H2-title', wrap=False)

def corner_label(slide, text):
    return tb(slide, 730, 25, 200, 18, [P1(text, 12, BODY)], align='r', name='corner-label', wrap=False)

def credit(slide, text):
    return tb(slide, 33, 511, 600, 12, [P1(text, 8, AUX)], name='source-credit', wrap=False)

CHROME = ('H2-title', 'source-credit', 'corner-label')

# ----------------------------------------------------------------------------- colour & type remap
def _hex(v):
    return tuple(int(v[i:i + 2], 16) / 255 for i in (0, 2, 4))

def _lum(v):
    r, g, b = _hex(v)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b

def _sat(v):
    c = _hex(v)
    return (max(c) - min(c)) / (max(c) + 1e-6)

def ctx(el):
    p = el.getparent()
    while p is not None:
        t = etree.QName(p).localname
        if t in ('rPr', 'endParaRPr', 'defRPr'):
            return 'text'
        if t in ('ln', 'lnL', 'lnR', 'lnT', 'lnB'):
            return 'line'
        if t in ('spPr', 'grpSpPr', 'bgPr', 'tcPr'):
            return 'fill'
        p = p.getparent()
    return 'fill'

def auto_colour(v, kind, accent_on):
    """Fallback mapping for colours with no explicit rule: keep the light/dark hierarchy, drop the hue."""
    L, s = _lum(v), _sat(v)
    if s > 0.28 and 0.15 < L < 0.8:                      # a real accent colour
        return RED if accent_on else {'text': AUX, 'line': MIDL, 'fill': '9FA0A0'}[kind]
    if kind == 'text':
        return TITLE if L < 0.22 else BODY if L < 0.4 else AUX
    if kind == 'line':
        return MIDL if L < 0.45 else OUTL if L < 0.75 else HAIR
    if L > 0.93:
        return 'FFFFFF'
    g = int(round(L * 255))
    return f'{g:02X}{g:02X}{g:02X}'

def recolor(slide, accent_on, text_map=None, line_map=None, fill_map=None, accents=(), overrides=None, tag=''):
    """Map every srgbClr by context. Explicit maps first (per source deck), then `auto_colour`."""
    text_map, line_map, fill_map = text_map or {}, line_map or {}, fill_map or {}
    over = overrides or {}
    keep = {TITLE, BODY, AUX, HAIR, MIDL, OUTL, RED, REDL, 'FFFFFF', TOC_BG}
    for c in slide._element.iter(qn('a:srgbClr')):
        v = c.get('val').upper(); k = ctx(c)
        if v in over.get(k, {}):
            new = over[k][v]
        elif k == 'text' and v in text_map:
            new = text_map[v]
        elif k == 'line' and v in line_map:
            new = line_map[v]
        elif k == 'fill' and v in fill_map:
            new = fill_map[v]
        elif v in accents:
            new = RED if accent_on else {'text': AUX, 'line': MIDL, 'fill': '9FA0A0'}[k]
        elif v in keep:
            continue
        else:
            new = auto_colour(v, k, accent_on)
            log.append(f'{tag}: auto-mapped {k} {v} -> {new}')
        c.set('val', new)

def default_size_map(s, color_hex):
    """Source size (pt) -> layout size, for decks whose body sat at 15-18pt.
    Item headings stay at body size (11pt): they are set apart by colour (TITLE vs BODY) and position,
    which keeps dense research pages calm — the approved look. Statements/leads go to 14pt."""
    if s >= 23: return 14
    if s >= 15: return 11
    if s >= 13.5: return 10
    if s >= 12: return 9
    return 8

def retype(slide, size_map=default_size_map, table_size=10):
    """Re-size every run by role and apply the font pair; multi-line paragraphs get a 1.45-1.6 pitch."""
    for p in slide._element.iter(qn('a:p')):
        in_tbl = any(etree.QName(a).localname == 'tbl' for a in p.iterancestors())
        size = None
        for rp in list(p.iter(qn('a:rPr'))) + list(p.iter(qn('a:endParaRPr'))):
            sz = rp.get('sz')
            if sz:
                fill = rp.find(qn('a:solidFill'))
                clr = fill.find(qn('a:srgbClr')) if fill is not None else None
                col = clr.get('val').upper() if clr is not None else ''
                new = table_size if in_tbl else size_map(int(sz) / 100, col)
                rp.set('sz', str(int(new * 100))); size = size or new
            if in_tbl and rp.get('b'):
                rp.set('b', '0')
            set_fonts(rp)
        ln = p.find(qn('a:pPr') + '/' + qn('a:lnSpc'))
        if ln is not None and size:
            for ch in list(ln):
                ln.remove(ch)
            sp = etree.SubElement(ln, qn('a:spcPts'))
            sp.set('val', str(int(round(size * (1.6 if size <= 12 else 1.45) * 100))))

def scale_geometry(slide, k):
    """Scale every offset/extent (and table grid) by k — e.g. 0.8 for a 1200x675pt source."""
    for x in list(slide._element.iter(qn('a:xfrm'))) + list(slide._element.iter(qn('p:xfrm'))):
        for tag in ('a:off', 'a:chOff'):
            n = x.find(qn(tag))
            if n is not None:
                n.set('x', str(int(round(int(n.get('x')) * k)))); n.set('y', str(int(round(int(n.get('y')) * k))))
        for tag in ('a:ext', 'a:chExt'):
            n = x.find(qn(tag))
            if n is not None:
                n.set('cx', str(int(round(int(n.get('cx')) * k)))); n.set('cy', str(int(round(int(n.get('cy')) * k))))
    for g in slide._element.iter(qn('a:gridCol')):
        g.set('w', str(int(round(int(g.get('w')) * k))))
    for r in slide._element.iter(qn('a:tr')):
        r.set('h', str(int(round(int(r.get('h')) * k))))
    for tc in slide._element.iter(qn('a:tcPr')):
        for a in ('marL', 'marR', 'marT', 'marB'):
            if tc.get(a):
                tc.set(a, str(int(round(int(tc.get(a)) * k * 0.8))))

def restyle_tables(slide):
    """No fills; 0.5pt rules — darker under the header row; header text TITLE, body BODY."""
    for tbl in slide._element.iter(qn('a:tbl')):
        tp = tbl.find(qn('a:tblPr'))
        for a in ('firstRow', 'bandRow'):
            tp.set(a, '0')
        for ri, tr in enumerate(tbl.findall(qn('a:tr'))):
            for tc in tr.findall(qn('a:tc')):
                pr = tc.find(qn('a:tcPr'))
                if pr is None:
                    pr = etree.SubElement(tc, qn('a:tcPr'))
                for f in pr.findall(qn('a:solidFill')) + pr.findall(qn('a:noFill')):
                    pr.remove(f)
                nf = etree.Element(qn('a:noFill'))
                anchor = next((pr.find(qn(t)) for t in ('a:lnB', 'a:lnT', 'a:lnR', 'a:lnL') if pr.find(qn(t)) is not None), None)
                (anchor.addnext(nf) if anchor is not None else pr.insert(0, nf))
                lnB = pr.find(qn('a:lnB'))
                if lnB is not None and lnB.find(qn('a:solidFill')) is not None:
                    lnB.set('w', '6350')
                    lnB.find(qn('a:solidFill')).find(qn('a:srgbClr')).set('val', MIDL if ri == 0 else HAIR)
                for rp in tc.iter(qn('a:rPr')):
                    f = rp.find(qn('a:solidFill'))
                    if f is not None and f.find(qn('a:srgbClr')) is not None:
                        f.find(qn('a:srgbClr')).set('val', TITLE if ri == 0 else BODY)

def drop_chrome(slide, names=(), texts=()):
    """Remove running headers, page numbers, footers, background textures (by shape name or exact text)."""
    for el in top_shapes(slide):
        if name_of(el) in names or text_of(el).strip() in texts:
            remove(el)

def take_text(slide, name):
    el = find(slide, name)
    t = text_of(el).strip() if el is not None else ''
    remove(el)
    return t

def vfit(slide, top=74, bottom=490, skip=CHROME):
    """Centre the content block vertically in [top, bottom]; top-align if it is taller."""
    els = [el for el in top_shapes(slide) if name_of(el) not in skip]
    if not els:
        return
    bbs = [bbox(el) for el in els]
    y0 = min(b[1] for b in bbs); y1 = max(b[1] + b[3] for b in bbs)
    h = y1 - y0
    ny = top + max(0, (bottom - top - h) / 2) if h < bottom - top else top
    for el in els:
        shift(el, dy=ny - y0)

def content_bottom(slide, skip=CHROME):
    return max((bbox(e)[1] + bbox(e)[3]) for e in top_shapes(slide) if name_of(e) not in skip)

# ----------------------------------------------------------------------------- template builders (A, B, C, G)
def build_cover(slide, cn, en, topic_cn, topic_en, date, family_png):
    """A — cover / back cover. family_png: the cover image — a cut-out group of the works (products),
    or any hero image (a space, a brand board, a screen set, an artwork); it is fitted into the right half."""
    clear(slide); white_bg(slide)
    tb(slide, 40, 194, 280, 36, [P1(cn, 28, TITLE)], align='c', name='cover-title')
    tb(slide, 40, 236, 280, 30, [P1(en, 24, TITLE)], align='c', name='cover-title-en')
    if topic_cn:
        tb(slide, 40, 290, 280, 16, [P1(topic_cn, 12, BODY)], align='c', name='cover-topic')
    if topic_en:
        tb(slide, 40, 309, 280, 14, [P1(topic_en, 10, AUX)], align='c', name='cover-topic-en')
    tb(slide, 68, 462, 200, 16, [P1(date, 12, BODY)], name='cover-date')
    if family_png:
        im = Image.open(family_png)
        pic(slide, family_png, *fit(im.width / im.height, 440, 150, 472, 310), name='cover-family')

def build_toc(slide, rows, art_png=None, art_box=(600, 92, None, 372)):
    """B — contents. rows: [(cn, en)]; numbered P/01..; optional grey line-art on the right."""
    clear(slide); white_bg(slide, TOC_BG)
    pitch = 40
    y0 = 270 - (len(rows) * pitch) / 2 + 8
    enx = max(176, 66 + 14 * max(len(cn) for cn, _ in rows) + 14)   # one EN column for every row
    for k, (cn, en) in enumerate(rows):
        y = y0 + k * pitch
        tb(slide, 66, y, 14 * len(cn) + 10, 20, [P1(cn, 14, TITLE)], name=f'toc-cn-{k}', wrap=False)
        tb(slide, enx, y + 4, 170, 18, [P1(en, 12, BODY)], name=f'toc-en-{k}', wrap=False)
        line(slide, enx + 169, y + 13, enx + 269, y + 13, HAIR, 0.5, name=f'toc-rule-{k}')
        tb(slide, enx + 286, y + 7, 40, 12, [P1(f'P/{k + 1:02d}', 8, AUX)], name=f'toc-page-{k}', wrap=False)
    if art_png:
        x, y, w, h = art_box
        pic(slide, art_png, x, y, w=w, h=h, name='toc-image')

def build_divider(slide, cn, en, art_png, art_box, tagline=None, tagline_en=None):
    """C — chapter divider: CN 28 / EN 26 centred in the upper third, one illustration below."""
    clear(slide); white_bg(slide)
    tb(slide, 180, 116, 600, 38, [P1(cn, 28, TITLE)], align='c', name='divider-title')
    tb(slide, 180, 162, 600, 34, [P1(en, 26, TITLE)], align='c', name='divider-title-en')
    if tagline:
        tb(slide, 180, 212, 600, 18, [P1(tagline, 12, AUX)], align='c', name='divider-tagline')
    if tagline_en:
        tb(slide, 180, 232, 600, 14, [P1(tagline_en, 10, AUX)], align='c', name='divider-tagline-en')
    if art_png:
        im = Image.open(art_png)
        pic(slide, art_png, *fit(im.width / im.height, *art_box), name='divider-art')

def build_g(slide, en, cn, art_png=None, art_box=(160, 150, 640, 330)):
    """G — unit title page (one work / product / space / series / chapter of a project):
    EN 28 centred, CN 12 below, a teaser image centred (motion-blurred cut-out, a crop, or a scene)."""
    white_bg(slide)
    tb(slide, 180, 76, 600, 36, [P1(en, 28, TITLE)], align='c', name='product-title-en')
    if cn:
        tb(slide, 180, 116, 600, 16, [P1(cn, 12, AUX)], align='c', name='product-title-cn')
    if art_png:
        im = Image.open(art_png)
        pic(slide, art_png, *fit(im.width / im.height, *art_box), name='product-hero')

def build_views(slide, items, top=74, bottom=490, max_h=320, gap=48, label_h=26, label_gap=10):
    """Views page (I variant): several cut-out angles of one product at equal height on one baseline,
    each labelled CN 10 / EN 8 underneath. items: [(png, cn_label, en_label)]."""
    ims = [Image.open(p) for p, _, _ in items]
    asp = [im.width / im.height for im in ims]
    n = len(items)
    h = min(max_h, bottom - top - label_h - label_gap, (840 - gap * (n - 1)) / sum(asp))
    x = 480 - (h * sum(asp) + gap * (n - 1)) / 2
    y = top + (bottom - top - (h + label_gap + label_h)) / 2
    for (p, cn, en), a in zip(items, asp):
        w = h * a
        pic(slide, p, x, y, w=w, h=h, name='view-image')
        tb(slide, x - 30, y + h + label_gap, w + 60, label_h,
           [P1(cn, 10, BODY), P1(en, 8, AUX, before=1)], align='c', name='view-label')
        x += w + gap

def set_lines(el, lines):
    """Replace the text of a (single-paragraph) text box with new lines, keeping the first run's formatting."""
    txb = txbody(el)
    paras = txb.findall(qn('a:p'))
    p0 = paras[0]
    for p in paras[1:]:
        txb.remove(p)
    r0 = p0.find(qn('a:r'))
    rPr = r0.find(qn('a:rPr')) if r0 is not None else None
    for ch in list(p0):
        if etree.QName(ch).localname in ('r', 'br', 'fld'):
            p0.remove(ch)
    end = p0.find(qn('a:endParaRPr'))
    for k, text in enumerate(lines):
        if k:
            br = etree.Element(qn('a:br'))
            if rPr is not None:
                br.append(copy.deepcopy(rPr))
            (end.addprevious(br) if end is not None else p0.append(br))
        r = etree.Element(qn('a:r'))
        if rPr is not None:
            r.append(copy.deepcopy(rPr))
        etree.SubElement(r, qn('a:t')).text = text
        (end.addprevious(r) if end is not None else p0.append(r))

def set_region(slide, el, path, cx, cy, w):
    """Show a region of a new image in an existing picture frame: centre (cx, cy) and width w in image pixels;
    the height follows the frame's aspect. Region is shifted to stay inside the image."""
    replace_blob(slide, el, path)
    W, H = Image.open(path).size
    fx, fy, fw, fh = bbox(el)
    w = min(w, W); h = min(w * fh / fw, H); w = h * fw / fh
    x0 = min(max(0, cx - w / 2), W - w); y0 = min(max(0, cy - h / 2), H - h)
    p = first_pic(el)
    bf = p.find(qn('p:blipFill'))
    sr = bf.find(qn('a:srcRect'))
    if sr is None:
        sr = etree.Element(qn('a:srcRect'))
        bf.find(qn('a:blip')).addnext(sr)
    for k, v in (('l', x0 / W), ('t', y0 / H), ('r', 1 - (x0 + w) / W), ('b', 1 - (y0 + h) / H)):
        sr.set(k, str(int(round(v * 100000))))

# ----------------------------------------------------------------------------- fill-in slots (待补图版)
SLOT_FILL, SLOT_LINE = 'F0F0EE', 'B5B5B6'
SLOT_STEPS = '① 复制上面的提示词去生成　② 删除本提示框　③ 右键灰框 →「更改图片」'
OLD_STEPS = '① 复制提示词去生成　② 删除本提示框　③ 右键下面的旧图 →「更改图片」'

def slot_png(path, aspect, label, font_file=None, long_px=900):
    """Light grey frame with a dashed border and a faint centred label, at the slot's aspect.
    Replacing it with 更改图片 keeps the frame's position and size."""
    from PIL import ImageDraw
    w, h = (long_px, max(2, round(long_px / aspect))) if aspect >= 1 else (max(2, round(long_px * aspect)), long_px)
    im = Image.new('RGB', (w, h), tuple(int(SLOT_FILL[i:i + 2], 16) for i in (0, 2, 4)))
    d = ImageDraw.Draw(im)
    line = tuple(int(SLOT_LINE[i:i + 2], 16) for i in (0, 2, 4))
    dash, gap, t = 14, 9, 3
    for x0 in range(0, w, dash + gap):
        d.line([(x0, t // 2), (min(w, x0 + dash), t // 2)], fill=line, width=t)
        d.line([(x0, h - 1 - t // 2), (min(w, x0 + dash), h - 1 - t // 2)], fill=line, width=t)
    for y0 in range(0, h, dash + gap):
        d.line([(t // 2, y0), (t // 2, min(h, y0 + dash))], fill=line, width=t)
        d.line([(w - 1 - t // 2, y0), (w - 1 - t // 2, min(h, y0 + dash))], fill=line, width=t)
    try:
        f = ImageFont.truetype(font_file, max(14, min(w, h) // 12)) if font_file else ImageFont.load_default()
    except Exception:
        f = ImageFont.load_default()
    bb = d.textbbox((0, 0), label, font=f)                     # bottom of the frame: the prompt card sits on top
    d.text(((w - bb[2] + bb[0]) / 2, h - (bb[3] - bb[1]) - h * 0.07 - bb[1]), label, font=f, fill=(201, 202, 202))
    im.save(path)
    return path

def _fit_card(slide, x, y, w, h, head, ref, prompt, steps, name, sizes=(8, 7.5, 7, 6.5, 6, 5.5, 5)):
    """Text box whose prompt shrinks until the whole card fits (w, h)."""
    for s in sizes:
        paras = [{'runs': [(head, s + 1, TITLE)], 'ln': (s + 1) * 1.3}]
        if ref:
            paras.append({'runs': [(ref, s, AUX)], 'ln': s * 1.3, 'before': 2})
        paras.append({'runs': [(prompt, s, BODY)], 'ln': s * 1.28, 'before': 4})
        paras.append({'runs': [(steps, max(5, s - 0.5), RED)], 'ln': max(5, s - 0.5) * 1.3, 'before': 4})
        shp = tb(slide, x, y, w, h, paras, name=name)
        if measure(shp.text_frame._txBody, w - 2) <= h or s == sizes[-1]:
            if s == sizes[-1] and measure(shp.text_frame._txBody, w - 2) > h:
                log.append(f'card {name}: prompt does not fit at {s}pt')
            return shp
        remove(shp._element)

SLOT_STEPS_SHORT = '复制提示词生成 → 删本框 → 右键灰框「更改图片」'

def add_slot(slide, x, y, w, h, sid, head, ref, prompt, workdir, font_file=None, pad=8, card=True, steps=None):
    """AI-image slot: grey placeholder picture at (x,y,w,h) + a prompt card inside it.
    Returns the picture. The prompt is also appended to the slide notes. Narrow slots get the short steps line."""
    steps = steps or (SLOT_STEPS if w >= 240 else SLOT_STEPS_SHORT)
    os.makedirs(workdir, exist_ok=True)
    ratio = w / h
    png = slot_png(os.path.join(workdir, f'slot_{sid}.png'), ratio, f'{sid}  AI 图片位', font_file)
    p = pic(slide, png, x, y, w=w, h=h, name=f'待替换图片-{sid}')
    if card:
        _fit_card(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, f'【{sid}】{head}', ref, prompt, steps,
                  name=f'提示词-替换后删除-{sid}')
    add_notes(slide, f'【{sid}】{head}\n{ref}\n{prompt}')
    return p

def add_old_card(slide, x, y, w, h, sid, head, ref, prompt, steps=None):
    """Prompt card over an existing (outdated) picture: white, slightly translucent, thin accent-colour border."""
    shp = _fit_card(slide, x, y, w, h, f'【{sid}】{head}', ref, prompt, steps or OLD_STEPS, name=f'提示词-替换后删除-{sid}')
    bp = shp.text_frame._txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, str(int(7 * P)))
    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    clr = shp._element.spPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
    etree.SubElement(clr, qn('a:alpha')).set('val', '93000')
    shp.line.color.rgb = RGBColor.from_string(RED); shp.line.width = Pt(0.75)
    add_notes(slide, f'【{sid}】{head}\n{ref}\n{prompt}')
    return shp

def add_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = (tf.text.rstrip() + '\n\n' if tf.text.strip() else '') + text

def text_block(slide, els, x, center_y, gaps, width=260):
    """Stack existing text shapes at x, centred on center_y. gaps[k] = space before element k."""
    hs = [bbox(e)[3] for e in els]
    y = center_y - (sum(hs) + sum(gaps[1:])) / 2
    for k, e in enumerate(els):
        if k:
            y += gaps[k]
        b = bbox(e)
        place(e, x, y, max(b[2], width), b[3])
        y += b[3]

# ----------------------------------------------------------------------------- bilingual copy
EN_SIZE = {14: 11, 13: 10, 12: 10, 11: 9, 10: 8.5, 9: 8}

def _font():
    if 'f' not in _FONT:
        path = CFG['en_font_file']
        _FONT['f'] = ImageFont.truetype(path, 200, index=0) if path and os.path.exists(path) else None
    return _FONT['f']

def full_width(ch):
    return '⺀' <= ch <= '￯' or ch in '“”‘’'

def text_w(s, size):
    f = _font()
    return sum(size if full_width(ch) else (f.getlength(ch) * size / 200 if f else 0.46 * size) for ch in s)

def wrap_count(s, size, width):
    if not s.strip():
        return 1
    if any(full_width(c) for c in s) and ' ' not in s.strip():
        return max(1, math.ceil(text_w(s, size) / width - 0.02))
    lines, cur, sp = 1, 0, text_w(' ', size)
    for w in s.split(' '):
        ww = text_w(w, size)
        if cur == 0:
            cur = ww
        elif cur + sp + ww <= width:
            cur += sp + ww
        else:
            lines += 1; cur = ww
    return lines

def para_info(p):
    rp = p.find('.//' + qn('a:rPr'))
    size = int(rp.get('sz')) / 100 if rp is not None and rp.get('sz') else 11
    pPr = p.find(qn('a:pPr'))
    pitch, before = size * 1.2, 0
    if pPr is not None:
        sp = pPr.find(qn('a:lnSpc') + '/' + qn('a:spcPts'))
        if sp is not None:
            pitch = int(sp.get('val')) / 100
        sb = pPr.find(qn('a:spcBef') + '/' + qn('a:spcPts'))
        if sb is not None:
            before = int(sb.get('val')) / 100
    segs, cur = [], ''
    for ch in p:
        t = etree.QName(ch).localname
        if t == 'r':
            cur += ''.join(x.text or '' for x in ch.iter(qn('a:t')))
        elif t == 'br':
            segs.append(cur); cur = ''
    segs.append(cur)
    return size, pitch, before, segs

def measure(txb, width):
    h = 0
    for k, p in enumerate(txb.findall(qn('a:p'))):
        size, pitch, before, segs = para_info(p)
        h += (before if k else 0) + pitch * sum(wrap_count(s, size, width) for s in segs)
    return h

def txbody(el):
    sp = el if etree.QName(el).localname == 'sp' else el.find('.//' + qn('p:sp'))
    return sp.find(qn('p:txBody'))

def fit_height(el, width=400):
    """Shrink a text box's height to its measured content (boxes keep their pre-restyle heights otherwise)."""
    b = bbox(el)
    place(el, b[0], b[1], b[2], measure(txbody(el), width) + 1)

def set_pitch(p, pitch, before=None):
    pPr = p.find(qn('a:pPr'))
    if pPr is None:
        pPr = etree.Element(qn('a:pPr')); p.insert(0, pPr)
    for tag in ('a:lnSpc', 'a:spcBef'):
        for n in pPr.findall(qn(tag)):
            pPr.remove(n)
    ln = etree.Element(qn('a:lnSpc'))
    etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(round(pitch * 100))))
    pPr.insert(0, ln)
    if before:
        sb = etree.Element(qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', str(int(round(before * 100))))
        ln.addnext(sb)
    return pPr

def en_run(text, size, color):
    r = etree.Element(qn('a:r'))
    rPr = etree.SubElement(r, qn('a:rPr'))
    rPr.set('lang', 'en-US'); rPr.set('sz', str(int(round(size * 100)))); rPr.set('dirty', '0')
    f = etree.SubElement(rPr, qn('a:solidFill'))
    etree.SubElement(f, qn('a:srgbClr')).set('val', color)
    set_fonts(rPr)
    etree.SubElement(r, qn('a:t')).text = text
    return r

def add_en(slide, el, spec, width=None):
    """Append the English counterpart inside the same text box and resize it. Returns (old_h, new_h).
    spec = ('block', 'text\\nnext line') | ('pairs', [en_for_cn_para_1, ...]) | ('inline', 'text')."""
    txb = txbody(el)
    paras = txb.findall(qn('a:p'))
    rp = txb.find('.//' + qn('a:rPr'))
    cn = int(rp.get('sz')) / 100
    fill = rp.find(qn('a:solidFill'))
    col = fill.find(qn('a:srgbClr')).get('val') if fill is not None else BODY
    ens = EN_SIZE.get(round(cn), round(cn * 1.6) / 2)
    en_col = AUX if col == AUX else BODY
    pPr0 = paras[0].find(qn('a:pPr'))
    algn = pPr0.get('algn', 'l') if pPr0 is not None else 'l'
    cn_pitch = cn * (1.55 if cn <= 12 else 1.4)
    en_pitch = ens * 1.42
    x, y, w, h = bbox(el)
    cn_w = max(text_w(s, cn) for p in paras for s in para_info(p)[3]) + 2
    for p in paras:
        set_pitch(p, cn_pitch).set('algn', algn)
    mode, payload = spec
    if mode == 'inline':
        end = paras[0].find(qn('a:endParaRPr'))
        r = en_run('  ' + payload, ens, en_col)
        (end.addprevious(r) if end is not None else paras[0].append(r))
        W = cn_w + text_w('  ' + payload, ens) + 4
    else:
        if mode == 'block':
            for k, ln_text in enumerate(payload.split('\n')):
                p = etree.SubElement(txb, qn('a:p'))
                set_pitch(p, en_pitch, (4 if cn <= 12 else 5) if k == 0 else None).set('algn', algn)
                p.append(en_run(ln_text, ens, en_col))
        else:
            for k, (p, ln_text) in enumerate(zip(paras, payload)):
                if k:
                    set_pitch(p, cn_pitch, 8).set('algn', algn)
                n = etree.Element(qn('a:p'))
                set_pitch(n, en_pitch, 2).set('algn', algn)
                n.append(en_run(ln_text, ens, en_col))
                p.addnext(n)
        W = max(cn_w, width or 0)
    txb.find(qn('a:bodyPr')).set('wrap', 'square')
    H = measure(txb, W - 2) + 1
    nx = x + (w - W) / 2 if algn == 'ctr' else x + w - W if algn == 'r' else x
    place(el, nx, y, W, H)
    return h, H

def avail_width(slide, el, cap=250, extra=70):
    """How wide an EN paragraph may run before it meets the next shape to the right."""
    x, y, w, h = bbox(el)
    txb = txbody(el)
    rp = txb.find('.//' + qn('a:rPr')); cn = int(rp.get('sz')) / 100
    cn_w = max(text_w(s, cn) for p in txb.findall(qn('a:p')) for s in para_info(p)[3]) + 2
    right = 900
    for o in top_shapes(slide):
        if o is el or name_of(o) in CHROME:
            continue
        ox, oy, ow, oh = bbox(o)
        if ox >= x + cn_w - 1 and oy < y + h + extra and oy + oh > y:
            right = min(right, ox - 14)
    return min(cap, right - x)

def reflow(slide, grown, rows=False):
    """Push shapes down below boxes that grew. Column mode moves only shapes overlapping in x
    (independent columns stay put); row mode moves everything below (row-based lists)."""
    els = [e for e in top_shapes(slide) if name_of(e) not in CHROME]
    geo = {}
    for e in els:
        x, y, w, h = bbox(e)
        oh, nh = grown.get(e, (h, h))
        geo[e] = (x, y, w, oh, nh - oh)
    order = sorted(els, key=lambda e: geo[e][1])
    shift_by = {}
    for e in order:
        x, y, w, oh, d = geo[e]
        s = 0
        for o in order:
            if o is e or o not in shift_by:
                continue
            ox, oy, ow, ooh, od = geo[o]
            if oy + ooh > y + 2:
                continue
            if not rows and not (ox < x + w - 3 and x < ox + ow - 3):
                continue
            s = max(s, shift_by[o] + od)
        shift_by[e] = s
    for e, s in shift_by.items():
        if s:
            shift(e, dy=s)

def apply_en(slide, specs, rows=False, tag=''):
    """specs: {shape_name: spec}. Adds EN to each, then reflows the page."""
    grown = {}
    for nm, spec in specs.items():
        el = find(slide, nm)
        if el is None:
            log.append(f'{tag}: EN target {nm} not found'); continue
        grown[el] = add_en(slide, el, spec, width=avail_width(slide, el))
    if grown:
        reflow(slide, grown, rows=rows)
    return bool(grown)

def anchor(slide, a, b, mode):
    """Re-attach shape a to shape b after a reflow. mode: 'c' centre on b, ('below', gap), or a y offset."""
    ea, eb = find(slide, a), find(slide, b)
    if mode == 'c':
        move_to(ea, y=bbox(eb)[1] + bbox(eb)[3] / 2 - bbox(ea)[3] / 2)
    elif isinstance(mode, tuple):
        move_to(ea, y=bbox(eb)[1] + bbox(eb)[3] + mode[1])
    else:
        move_to(ea, y=bbox(eb)[1] + mode)

# ----------------------------------------------------------------------------- deck-level
def blank_deck(base=None):
    """Start a deck from scratch at 960x540. Prefer base= an existing 16:9 .pptx whose fonts this machine has
    (e.g. the user's own file): all its slides are removed, its masters and theme are kept.
    python-pptx's built-in default template can make Keynote stop on a dialog when rendering."""
    from pptx import Presentation
    prs = Presentation(base) if base else Presentation()
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.get(qn('r:id')))
        lst.remove(sid)
    prs.slide_width, prs.slide_height = E(PAGE_W), E(PAGE_H)
    return prs

def new_slide(prs, layout_name='Blank'):
    lay = next((l for l in prs.slide_layouts if l.name == layout_name), prs.slide_layouts[-1])
    s = prs.slides.add_slide(lay)
    for ph in list(s.placeholders):
        remove(ph._element)
    return s

def reorder(prs, order):
    """order: list of slide objects (every slide exactly once)."""
    lst = prs.slides._sldIdLst
    els = [next(el for el in lst if int(el.get('id')) == s.slide_id) for s in order]
    for el in list(lst):
        lst.remove(el)
    for el in els:
        lst.append(el)

def set_theme_fonts(pptx_path):
    """Write the font pair into the theme so text the user adds later matches."""
    tmp = pptx_path + '.tmp'
    with zipfile.ZipFile(pptx_path) as zi, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if re.match(r'ppt/theme/theme\d+\.xml$', it.filename):
                x = data.decode('utf8')
                x = re.sub(r'<a:latin typeface="[^"]*"/>', f'<a:latin typeface="{CFG["en_font"]}"/>', x)
                x = re.sub(r'<a:ea typeface="[^"]*"/>', f'<a:ea typeface="{CFG["cn_font"]}"/>', x)
                x = re.sub(r'<a:cs typeface="[^"]*"/>', f'<a:cs typeface="{CFG["cn_font"]}"/>', x)
                data = x.encode('utf8')
            zo.writestr(it, data)
    shutil.move(tmp, pptx_path)
